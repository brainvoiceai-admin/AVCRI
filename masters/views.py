from django.http import JsonResponse, HttpResponse
from django.views.decorators.http import require_http_methods
from django.views.decorators.csrf import csrf_exempt
from django.contrib.auth.decorators import login_required
from django.shortcuts import render, get_object_or_404, redirect
from django.contrib import messages
from django.views.generic import ListView, CreateView, UpdateView, DeleteView
from django.urls import reverse_lazy
from django.utils.decorators import method_decorator
from django.db import models as db_models
from django import forms
from authentication.decorators import super_admin_required, module_access_required
from .models import (
    Procured, OilClass, UsageType, PrescriptionUsage, PrescriptionTime, 
    PrescriptionScope, UnitNo, TherapyMaster, TherapyGroup, ItemMaster, ItemGroup, SlotMaster,
    TherapistGrade, TherapistMaster, TherapistLeaveMaster, NurseTreatmentMaster,
    NurseMaster, IngredientMaster,
    ShiftMaster, ShiftSlotMapping, TherapistShiftMapping
)
from Patient.models import TreatmentRoom, RoomName, Floor
import json
import pandas as pd
from datetime import datetime
from io import BytesIO
from django.db import transaction, connection
from django.db.models import Max


def _log_master_audit(request, action, object_type, object_id, description, additional_data=None):
    """Log master CRUD to audit log (authentication.AuditLog)."""
    try:
        from authentication.views import log_audit_event
        if request and getattr(request, 'user', None) and getattr(request.user, 'is_authenticated', False):
            log_audit_event(
                user=request.user,
                action=action,
                description=description,
                status='success',
                object_type=object_type,
                object_id=str(object_id),
                additional_data=additional_data,
                request=request
            )
    except Exception:
        pass


def generate_unique_code(prefix, model_class, field_name):
    """
    Generates a unique code in the format PREFIX + 3 digits (e.g., ITM001).
    Ensures no collision and monotonically increasing numbers.
    """
    with transaction.atomic():
        # Get all existing codes
        existing_codes = model_class.objects.filter(
            **{f"{field_name}__startswith": prefix}
        ).values_list(field_name, flat=True)
        
        max_number = 0
        for code in existing_codes:
            try:
                # Extract numeric part
                # Assuming code format is exactly PREFIX + Number
                number_part = code[len(prefix):]
                if number_part.isdigit():
                    num = int(number_part)
                    if num > max_number:
                        max_number = num
            except Exception:
                continue
        
        next_number = max_number + 1
        return f"{prefix}{next_number:03d}"

def generate_numeric_code(model_class, field_name, start_from=1000):
    """
    Generates a unique numeric code (4 to 5 digits, e.g., '1001').
    """
    with transaction.atomic():
        existing_codes = model_class.objects.values_list(field_name, flat=True)
        max_number = 0
        for code in existing_codes:
            try:
                # Try to parse as int
                num = int(str(code))
                if num > max_number:
                    max_number = num
            except ValueError:
                continue
        
        if max_number < start_from:
            next_number = start_from
        else:
            next_number = max_number + 1
            
        return str(next_number)


@login_required
@require_http_methods(["GET"])
def get_master_data(request):
    """API endpoint to get master data for dropdowns"""
    master_type = request.GET.get('type', '')
    
    try:
        if master_type == 'therapy':
            data = list(TherapyMaster.objects.filter(is_active=True).values(
                'therapy_code', 'therapy_name', 'therapy_group__name', 'therapy_category'
            ))
        elif master_type == 'therapy_group':
            data = list(TherapyGroup.objects.filter(is_active=True).values('id', 'name'))
        elif master_type == 'item':
            data = list(ItemMaster.objects.filter(is_active=True).values(
                'item_code', 'item_name', 'item_group__name'
            ))
        elif master_type == 'item_group':
            data = list(ItemGroup.objects.filter(is_active=True).values('id', 'name'))
        elif master_type == 'slot':
            data = list(SlotMaster.objects.filter(is_active=True).values(
                'slot_number', 'slot_name', 'slot_time'
            ))
        elif master_type == 'procured':
            data = list(Procured.objects.filter(is_active=True).values('id', 'name'))
        elif master_type == 'oil_class':
            data = list(OilClass.objects.filter(is_active=True).values('id', 'name'))
        elif master_type == 'usage_type':
            data = list(UsageType.objects.filter(is_active=True).values('id', 'name'))
        elif master_type == 'prescription_time':
            data = list(PrescriptionTime.objects.filter(is_active=True).values('id', 'name'))
        elif master_type == 'prescription_scope':
            data = list(PrescriptionScope.objects.filter(is_active=True).values('id', 'name'))
        elif master_type == 'unit_no':
            data = list(UnitNo.objects.filter(is_active=True).values('id', 'unit_name'))
        elif master_type == 'therapist':
            data = list(TherapistMaster.objects.filter(is_active=True).values(
                'therapist_code', 'therapist_name', 'therapist_grade', 'employee_code', 'is_bathing_therapist'
            ))
        elif master_type == 'therapist_leave':
            therapist_code = request.GET.get('therapist_code', '')
            if therapist_code:
                data = list(TherapistLeaveMaster.objects.filter(
                    therapist__therapist_code=therapist_code, is_active=True
                ).values('leave_day', 'leave_day_name', 'is_active'))
            else:
                data = list(TherapistLeaveMaster.objects.filter(is_active=True).values(
                    'therapist__therapist_code', 'therapist__therapist_name', 'leave_day', 'leave_day_name'
                ))
        else:
            return JsonResponse({'error': 'Invalid master type'}, status=400)
        
        return JsonResponse({'data': data})
        
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)

@login_required
@super_admin_required
@require_http_methods(["POST"])
@csrf_exempt
def create_master_item(request):
    """API endpoint to create master items (Super Admin only)"""
    try:
        data = json.loads(request.body)
        master_type = data.get('type', '')
        
        if master_type == 'therapy':
            # Validate required fields (all except therapy_group)
            required = ['therapy_code', 'therapy_name', 'therapy_category', 'therapy_duration',
                        'bathing_required', 'treatment_room_required', 'cost_of_therapy']
            missing = []
            for f in required:
                val = data.get(f)
                if f == 'cost_of_therapy':
                    if val is None or val == '':
                        missing.append(f)
                elif not val:
                    missing.append(f)
            if missing:
                return JsonResponse({'error': f'Missing required fields: {", ".join(missing)}'}, status=400)
            # Get therapy_group if provided (optional)
            therapy_group = None
            if data.get('therapy_group'):
                try:
                    if isinstance(data.get('therapy_group'), int):
                        therapy_group = TherapyGroup.objects.get(id=data.get('therapy_group'))
                    else:
                        therapy_group = TherapyGroup.objects.get(name=data.get('therapy_group'))
                except TherapyGroup.DoesNotExist:
                    pass
            
            item = TherapyMaster.objects.create(
                therapy_code=data['therapy_code'],
                therapy_name=data['therapy_name'],
                therapy_group=therapy_group,
                therapy_category=data['therapy_category'],
                therapy_duration=data['therapy_duration'],
                bathing_required=data['bathing_required'],
                treatment_room_required=data['treatment_room_required'],
                cost_of_therapy=data.get('cost_of_therapy', 0) or 0
            )
        elif master_type == 'item':
            # Get item_group if provided
            item_group = None
            if data.get('item_group'):
                try:
                    if isinstance(data.get('item_group'), int):
                        item_group = ItemGroup.objects.get(id=data.get('item_group'))
                    else:
                        item_group = ItemGroup.objects.get(name=data.get('item_group'))
                except ItemGroup.DoesNotExist:
                    pass
            
            # Get umo if provided
            umo = None
            if data.get('umo'):
                try:
                    if isinstance(data.get('umo'), int):
                        umo = UnitNo.objects.get(id=data.get('umo'))
                    else:
                        umo = UnitNo.objects.get(unit_name=data.get('umo'))
                except UnitNo.DoesNotExist:
                    pass
            
            item = ItemMaster.objects.create(
                item_name=data['item_name'],
                item_group=item_group,
                umo=umo,
                min_quantity=data.get('min_quantity'),
                max_quantity=data.get('max_quantity')
            )
        elif master_type == 'slot':
            from datetime import datetime
            slot_time_str = data.get('slot_time') or data.get('start_time', '09:00')
            # Convert string to time if needed
            if isinstance(slot_time_str, str):
                try:
                    slot_time = datetime.strptime(slot_time_str, '%H:%M').time()
                except:
                    slot_time = datetime.strptime(slot_time_str, '%H:%M:%S').time()
            else:
                slot_time = slot_time_str
            item = SlotMaster.objects.create(
                slot_number=data['slot_number'],
                slot_time=slot_time,
                slot_name=data.get('slot_name', '')
            )
        elif master_type == 'ingredient':
            item = IngredientMaster.objects.create(
                ingredient_code=data['ingredient_code'],
                ingredient_name=data['ingredient_name'],
                description=data.get('description', ''),
                is_active=data.get('is_active', True)
            )
        else:
            return JsonResponse({'error': 'Invalid master type'}, status=400)
        
        obj_name = getattr(item, 'therapy_code', None) or getattr(item, 'item_name', None) or getattr(item, 'slot_number', None) or getattr(item, 'ingredient_code', None) or str(item.id)
        _log_master_audit(request, 'create', master_type.title(), str(item.id), f'Created {master_type}: {obj_name}', {'master_type': master_type})
        return JsonResponse({'success': True, 'id': item.id})
        
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)

@login_required
@super_admin_required
@require_http_methods(["PUT"])
@csrf_exempt
def update_master_item(request, item_id):
    """API endpoint to update master items (Super Admin only)"""
    try:
        data = json.loads(request.body)
        master_type = data.get('type', '')
        
        if master_type == 'therapy':
            item = TherapyMaster.objects.get(id=item_id)
            item.therapy_name = data.get('therapy_name', item.therapy_name)
            # Handle therapy_group as ForeignKey - can be ID or name
            therapy_group_data = data.get('therapy_group')
            if therapy_group_data:
                if isinstance(therapy_group_data, int):
                    try:
                        item.therapy_group = TherapyGroup.objects.get(id=therapy_group_data)
                    except TherapyGroup.DoesNotExist:
                        item.therapy_group = None
                elif isinstance(therapy_group_data, str):
                    try:
                        item.therapy_group = TherapyGroup.objects.get(name=therapy_group_data)
                    except TherapyGroup.DoesNotExist:
                        item.therapy_group = None
                else:
                    item.therapy_group = None
            else:
                item.therapy_group = None
            item.therapy_category = data.get('therapy_category', item.therapy_category)
            item.is_active = data.get('is_active', item.is_active)
        elif master_type == 'item':
            item = ItemMaster.objects.get(id=item_id)
            item.item_name = data.get('item_name', item.item_name)
            # Handle item_group as ForeignKey - can be ID or name
            item_group_data = data.get('item_group')
            if item_group_data:
                if isinstance(item_group_data, int):
                    try:
                        item.item_group = ItemGroup.objects.get(id=item_group_data)
                    except ItemGroup.DoesNotExist:
                        item.item_group = None
                elif isinstance(item_group_data, str):
                    try:
                        item.item_group = ItemGroup.objects.get(name=item_group_data)
                    except ItemGroup.DoesNotExist:
                        item.item_group = None
                else:
                    item.item_group = None
            else:
                item.item_group = None
            # Handle umo as ForeignKey - can be ID or unit_name
            umo_data = data.get('umo')
            if umo_data:
                if isinstance(umo_data, int):
                    try:
                        item.umo = UnitNo.objects.get(id=umo_data)
                    except UnitNo.DoesNotExist:
                        item.umo = None
                elif isinstance(umo_data, str):
                    try:
                        item.umo = UnitNo.objects.get(unit_name=umo_data)
                    except UnitNo.DoesNotExist:
                        item.umo = None
                else:
                    item.umo = None
            else:
                item.umo = None
            item.min_quantity = data.get('min_quantity', item.min_quantity)
            item.max_quantity = data.get('max_quantity', item.max_quantity)
            # Handle procured_from as ForeignKey - can be ID or name
            procured_from_data = data.get('procured_from')
            if procured_from_data:
                if isinstance(procured_from_data, int):
                    try:
                        item.procured_from = Procured.objects.get(id=procured_from_data)
                    except Procured.DoesNotExist:
                        item.procured_from = None
                elif isinstance(procured_from_data, str):
                    try:
                        item.procured_from = Procured.objects.get(name=procured_from_data)
                    except Procured.DoesNotExist:
                        item.procured_from = None
                else:
                    item.procured_from = None
            else:
                item.procured_from = None
            item.is_active = data.get('is_active', item.is_active)
        elif master_type == 'slot':
            item = SlotMaster.objects.get(id=item_id)
            item.slot_name = data.get('slot_name', item.slot_name)
            slot_time = data.get('slot_time') or data.get('start_time')
            if slot_time:
                from datetime import datetime
                if isinstance(slot_time, str):
                    try:
                        item.slot_time = datetime.strptime(slot_time, '%H:%M').time()
                    except:
                        item.slot_time = datetime.strptime(slot_time, '%H:%M:%S').time()
                else:
                    item.slot_time = slot_time
            item.is_active = data.get('is_active', item.is_active)
        elif master_type == 'ingredient':
            item = IngredientMaster.objects.get(id=item_id)
            item.ingredient_code = data.get('ingredient_code', item.ingredient_code)
            item.ingredient_name = data.get('ingredient_name', item.ingredient_name)
            item.description = data.get('description', item.description)
            item.is_active = data.get('is_active', item.is_active)
        else:
            return JsonResponse({'error': 'Invalid master type'}, status=400)
        
        item.save()
        obj_name = getattr(item, 'therapy_code', None) or getattr(item, 'therapy_name', None) or getattr(item, 'item_name', None) or getattr(item, 'slot_number', None) or getattr(item, 'ingredient_code', None) or str(item.id)
        _log_master_audit(request, 'update', master_type.title(), str(item_id), f'Updated {master_type}: {obj_name}', {'master_type': master_type})
        return JsonResponse({'success': True})
        
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)

@login_required
@super_admin_required
@require_http_methods(["DELETE"])
@csrf_exempt
def delete_master_item(request, item_id):
    """API endpoint to delete master items (Super Admin only)"""
    try:
        data = json.loads(request.body)
        master_type = data.get('type', '')
        
        if master_type == 'therapy':
            item = TherapyMaster.objects.get(id=item_id)
        elif master_type == 'item':
            item = ItemMaster.objects.get(id=item_id)
        elif master_type == 'slot':
            item = SlotMaster.objects.get(id=item_id)
        elif master_type == 'ingredient':
            item = IngredientMaster.objects.get(id=item_id)
        else:
            return JsonResponse({'error': 'Invalid master type'}, status=400)
        
        obj_name = getattr(item, 'therapy_code', None) or getattr(item, 'therapy_name', None) or getattr(item, 'item_name', None) or getattr(item, 'slot_number', None) or getattr(item, 'ingredient_code', None) or str(item_id)
        item.delete()
        _log_master_audit(request, 'delete', master_type.title(), str(item_id), f'Deleted {master_type}: {obj_name}', {'master_type': master_type})
        return JsonResponse({'success': True})
        
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)

# Class-based views for master management
class ProcuredManagementView(ListView):
    model = Procured
    template_name = 'masters/procured_manage.html'
    context_object_name = 'procured_list'
    paginate_by = 20

    def get_queryset(self):
        queryset = Procured.objects.all()
        
        # Search filter
        search = self.request.GET.get('search', '').strip()
        if search:
            queryset = queryset.filter(
                db_models.Q(name__icontains=search)
            )
        
        # Status filter
        status = self.request.GET.get('status', '').strip().lower()
        if status == 'active':
            queryset = queryset.filter(is_active=True)
        elif status == 'inactive':
            queryset = queryset.filter(is_active=False)
        # If status is empty, show all (no filter applied)
        
        return queryset.order_by('name')
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['master_type'] = 'procured'
        context['master_name'] = 'Procurement Sources'
        context['required_fields'] = ['name']
        context['optional_fields'] = ['is_active']
        context['search'] = self.request.GET.get('search', '')
        context['status_filter'] = self.request.GET.get('status', '')
        return context
    
    def post(self, request, *args, **kwargs):
        """Handle POST requests for creating new procurement sources"""
        name = request.POST.get('name')
        is_active = request.POST.get('is_active') == 'on'
        
        # Create new record
        if name:
            try:
                obj = Procured.objects.create(
                    name=name,
                    is_active=is_active
                )
                _log_master_audit(request, 'create', 'Procured', str(obj.id), f'Created Procured: {name}')
                messages.success(request, f'Procurement source "{name}" created successfully.')
            except Exception as e:
                messages.error(request, f'Error creating procurement source: {str(e)}')
        else:
            messages.error(request, 'Name is required.')
        
        return redirect('procured_manage')

class ProcuredUpdateView(UpdateView):
    model = Procured
    fields = ['name', 'is_active']
    template_name = 'masters/procured_form.html'
    success_url = reverse_lazy('procured_manage')
    
    def form_valid(self, form):
        _log_master_audit(self.request, 'update', 'Procured', str(self.object.pk), f'Updated Procured: {self.object.name}')
        messages.success(self.request, 'Procurement source updated successfully.')
        return super().form_valid(form)

@login_required
@super_admin_required
def procured_delete(request, pk):
    """Delete a procurement source - function-based view to handle GET requests properly"""
    procured = get_object_or_404(Procured, pk=pk)
    procured_name = procured.name
    procured.delete()
    _log_master_audit(request, 'delete', 'Procured', str(pk), f'Deleted Procured: {procured_name}')
    messages.success(request, f'Procurement source "{procured_name}" deleted successfully.')
    return redirect('procured_manage')

@login_required
@super_admin_required
def procured_enable(request, pk):
    """Enable a procurement source"""
    procured = get_object_or_404(Procured, pk=pk)
    procured.is_active = True
    procured.save()
    _log_master_audit(request, 'update', 'Procured', str(pk), f'Enabled Procured: {procured.name}')
    messages.success(request, f'Procurement source "{procured.name}" enabled successfully.')
    return redirect('procured_manage')

@login_required
@super_admin_required
def procured_disable(request, pk):
    """Disable a procurement source"""
    procured = get_object_or_404(Procured, pk=pk)
    procured.is_active = False
    procured.save()
    _log_master_audit(request, 'update', 'Procured', str(pk), f'Disabled Procured: {procured.name}')
    messages.success(request, f'Procurement source "{procured.name}" disabled successfully.')
    return redirect('procured_manage')

class OilClassManagementView(ListView):
    model = OilClass
    template_name = 'masters/oilclass_manage.html'
    context_object_name = 'oilclass_list'
    paginate_by = 20

    def get_queryset(self):
        # Get status filter FIRST
        status = self.request.GET.get('status', '').strip().lower()
        
        # Start with base queryset and apply status filter immediately
        if status == 'active':
            queryset = OilClass.objects.filter(is_active=True)
        elif status == 'inactive':
            queryset = OilClass.objects.filter(is_active=False)
        else:
            # No status filter - show all
            queryset = OilClass.objects.all()
        
        # Search filter
        search = self.request.GET.get('search', '').strip()
        if search:
            queryset = queryset.filter(name__icontains=search)
        
        return queryset.order_by('name')
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['master_type'] = 'oilclass'
        context['master_name'] = 'Oil Classes'
        context['required_fields'] = ['name']
        context['optional_fields'] = ['is_active']
        context['search'] = self.request.GET.get('search', '')
        context['status_filter'] = self.request.GET.get('status', '')
        return context
    
    def post(self, request, *args, **kwargs):
        """Handle POST requests for creating new oil classes"""
        name = request.POST.get('name')
        is_active = request.POST.get('is_active') == 'on'
        
        # Create new record
        if name:
            try:
                obj = OilClass.objects.create(name=name, is_active=is_active)
                _log_master_audit(request, 'create', 'OilClass', str(obj.id), f'Created OilClass: {name}')
                messages.success(request, f'Oil class "{name}" created successfully.')
            except Exception as e:
                messages.error(request, f'Error creating oil class: {str(e)}')
        else:
            messages.error(request, 'Name is required.')
        
        return redirect('oilclass_manage')

class OilClassUpdateView(UpdateView):
    model = OilClass
    fields = ['name', 'is_active']
    template_name = 'masters/oilclass_form.html'
    success_url = reverse_lazy('oilclass_manage')
    
    def form_valid(self, form):
        _log_master_audit(self.request, 'update', 'OilClass', str(self.object.pk), f'Updated OilClass: {self.object.name}')
        messages.success(self.request, 'Oil class updated successfully.')
        return super().form_valid(form)

@login_required
@super_admin_required
def oil_class_delete(request, pk):
    """Delete an oil class - function-based view to handle GET requests properly"""
    oil_class = get_object_or_404(OilClass, pk=pk)
    oil_class_name = oil_class.name
    oil_class.delete()
    _log_master_audit(request, 'delete', 'OilClass', str(pk), f'Deleted OilClass: {oil_class_name}')
    messages.success(request, f'Oil class "{oil_class_name}" deleted successfully.')
    return redirect('oilclass_manage')

@login_required
@super_admin_required
def oil_class_enable(request, pk):
    """Enable an oil class"""
    oil_class = get_object_or_404(OilClass, pk=pk)
    oil_class.is_active = True
    oil_class.save()
    _log_master_audit(request, 'update', 'OilClass', str(pk), f'Enabled OilClass: {oil_class.name}')
    messages.success(request, f'Oil class "{oil_class.name}" enabled successfully.')
    return redirect('oilclass_manage')

@login_required
@super_admin_required
def oil_class_disable(request, pk):
    """Disable an oil class"""
    oil_class = get_object_or_404(OilClass, pk=pk)
    oil_class.is_active = False
    oil_class.save()
    _log_master_audit(request, 'update', 'OilClass', str(pk), f'Disabled OilClass: {oil_class.name}')
    messages.success(request, f'Oil class "{oil_class.name}" disabled successfully.')
    return redirect('oilclass_manage')

class UsageTypeManagementView(ListView):
    model = UsageType
    template_name = 'masters/usagetype_manage.html'
    context_object_name = 'usagetype_list'
    paginate_by = 20

    def get_queryset(self):
        # Get status filter FIRST - ensure it's read correctly
        status_param = self.request.GET.get('status', '')
        if status_param:
            status = status_param.strip().lower()
        else:
            status = ''
        
        # Start with base queryset and apply status filter immediately
        # This ensures the filter is applied at the base queryset level
        if status == 'active':
            # Explicitly filter for active items only
            queryset = UsageType.objects.filter(is_active__exact=True)
        elif status == 'inactive':
            # Explicitly filter for inactive items only
            queryset = UsageType.objects.filter(is_active__exact=False)
        else:
            # No status filter - show all
            queryset = UsageType.objects.all()
        
        # Search filter - apply after status filter
        search = self.request.GET.get('search', '').strip()
        if search:
            queryset = queryset.filter(name__icontains=search)
        
        # Return ordered queryset
        return queryset.order_by('name')
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['master_type'] = 'usagetype'
        context['master_name'] = 'Usage Types'
        context['required_fields'] = ['name']
        context['optional_fields'] = ['is_active']
        context['search'] = self.request.GET.get('search', '')
        context['status_filter'] = self.request.GET.get('status', '')
        return context
    
    def post(self, request, *args, **kwargs):
        """Handle POST requests for creating new usage types"""
        name = request.POST.get('name')
        is_active = request.POST.get('is_active') == 'on'
        
        # Create new record
        if name:
            try:
                obj = UsageType.objects.create(name=name, is_active=is_active)
                _log_master_audit(request, 'create', 'UsageType', str(obj.id), f'Created UsageType: {name}')
                messages.success(request, f'Usage type "{name}" created successfully.')
            except Exception as e:
                messages.error(request, f'Error creating usage type: {str(e)}')
        else:
            messages.error(request, 'Name is required.')
        
        return redirect('usagetype_manage')

class UsageTypeUpdateView(UpdateView):
    model = UsageType
    fields = ['name', 'is_active']
    template_name = 'masters/usagetype_form.html'
    success_url = reverse_lazy('usagetype_manage')
    
    def form_valid(self, form):
        _log_master_audit(self.request, 'update', 'UsageType', str(self.object.pk), f'Updated UsageType: {self.object.name}')
        messages.success(self.request, 'Usage type updated successfully.')
        return super().form_valid(form)

@login_required
@super_admin_required
def usage_type_delete(request, pk):
    """Delete a usage type - function-based view to handle GET requests properly"""
    usage_type = get_object_or_404(UsageType, pk=pk)
    usage_type_name = usage_type.name
    usage_type.delete()
    _log_master_audit(request, 'delete', 'UsageType', str(pk), f'Deleted UsageType: {usage_type_name}')
    messages.success(request, f'Usage type "{usage_type_name}" deleted successfully.')
    return redirect('usagetype_manage')

@login_required
@super_admin_required
def usage_type_enable(request, pk):
    """Enable a usage type"""
    usage_type = get_object_or_404(UsageType, pk=pk)
    usage_type.is_active = True
    usage_type.save()
    _log_master_audit(request, 'update', 'UsageType', str(pk), f'Enabled UsageType: {usage_type.name}')
    messages.success(request, f'Usage type "{usage_type.name}" enabled successfully.')
    return redirect('usagetype_manage')

@login_required
@super_admin_required
def usage_type_disable(request, pk):
    """Disable a usage type"""
    usage_type = get_object_or_404(UsageType, pk=pk)
    usage_type.is_active = False
    usage_type.save()
    _log_master_audit(request, 'update', 'UsageType', str(pk), f'Disabled UsageType: {usage_type.name}')
    messages.success(request, f'Usage type "{usage_type.name}" disabled successfully.')
    return redirect('usagetype_manage')

# Prescription Usage Management Views
@method_decorator(module_access_required('prescription_usage_manage'), name='dispatch')
class PrescriptionUsageManagementView(ListView):
    model = PrescriptionUsage
    template_name = 'masters/prescriptionusage_manage.html'
    context_object_name = 'prescriptionusage_list'
    paginate_by = 20

    def get_queryset(self):
        queryset = PrescriptionUsage.objects.all()
        
        # Search filter
        search = self.request.GET.get('search', '')
        if search:
            queryset = queryset.filter(name__icontains=search)
        
        # Status filter
        status = self.request.GET.get('status', '')
        if status == 'active':
            queryset = queryset.filter(is_active=True)
        elif status == 'inactive':
            queryset = queryset.filter(is_active=False)
        
        return queryset.order_by('order_no', 'name')
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['master_type'] = 'prescriptionusage'
        context['master_name'] = 'Prescription Usage'
        context['required_fields'] = ['name']
        context['optional_fields'] = ['is_active']
        context['search'] = self.request.GET.get('search', '')
        context['status_filter'] = self.request.GET.get('status', '')
        return context
    
    def post(self, request, *args, **kwargs):
        """Handle POST requests for creating new prescription usages"""
        name = request.POST.get('name')
        order_no = request.POST.get('order_no', 0)
        try:
            order_no = int(order_no)
        except ValueError:
            order_no = 0
        is_active = request.POST.get('is_active') == 'on'
        
        # Create new record
        if name:
            try:
                obj = PrescriptionUsage.objects.create(name=name, order_no=order_no, is_active=is_active)
                _log_master_audit(request, 'create', 'PrescriptionUsage', str(obj.id), f'Created PrescriptionUsage: {name}')
                messages.success(request, f'Prescription usage "{name}" created successfully.')
            except Exception as e:
                messages.error(request, f'Error creating prescription usage: {str(e)}')
        else:
            messages.error(request, 'Name is required.')
        
        return redirect('prescriptionusage_manage')

@method_decorator(module_access_required('prescription_usage_manage'), name='dispatch')
class PrescriptionUsageUpdateView(UpdateView):
    model = PrescriptionUsage
    fields = ['name', 'order_no', 'is_active']
    template_name = 'masters/prescriptionusage_form.html'
    success_url = reverse_lazy('prescriptionusage_manage')
    
    def form_valid(self, form):
        _log_master_audit(self.request, 'update', 'PrescriptionUsage', str(self.object.pk), f'Updated PrescriptionUsage: {self.object.name}')
        messages.success(self.request, 'Prescription usage updated successfully.')
        return super().form_valid(form)

@login_required
@module_access_required('prescription_usage_manage')
def prescription_usage_delete(request, pk):
    """Delete a prescription usage - function-based view to handle GET requests properly"""
    prescription_usage = get_object_or_404(PrescriptionUsage, pk=pk)
    prescription_usage_name = prescription_usage.name
    prescription_usage.delete()
    _log_master_audit(request, 'delete', 'PrescriptionUsage', str(pk), f'Deleted PrescriptionUsage: {prescription_usage_name}')
    messages.success(request, f'Prescription usage "{prescription_usage_name}" deleted successfully.')
    return redirect('prescriptionusage_manage')

@login_required
@module_access_required('prescription_usage_manage')
def prescription_usage_enable(request, pk):
    """Enable a prescription usage"""
    prescription_usage = get_object_or_404(PrescriptionUsage, pk=pk)
    prescription_usage.is_active = True
    prescription_usage.save()
    _log_master_audit(request, 'update', 'PrescriptionUsage', str(pk), f'Enabled PrescriptionUsage: {prescription_usage.name}')
    messages.success(request, f'Prescription usage "{prescription_usage.name}" enabled successfully.')
    return redirect('prescriptionusage_manage')

@login_required
@module_access_required('prescription_usage_manage')
def prescription_usage_disable(request, pk):
    """Disable a prescription usage"""
    prescription_usage = get_object_or_404(PrescriptionUsage, pk=pk)
    prescription_usage.is_active = False
    prescription_usage.save()
    _log_master_audit(request, 'update', 'PrescriptionUsage', str(pk), f'Disabled PrescriptionUsage: {prescription_usage.name}')
    messages.success(request, f'Prescription usage "{prescription_usage.name}" disabled successfully.')
    return redirect('prescriptionusage_manage')

# Nurse Treatment Master Management Views
@method_decorator(super_admin_required, name='dispatch')
class NurseTreatmentMasterManagementView(ListView):
    model = NurseTreatmentMaster
    template_name = 'masters/nursetreatment_manage.html'
    context_object_name = 'nursetreatment_list'
    paginate_by = 20

    def get_queryset(self):
        queryset = NurseTreatmentMaster.objects.all()
        
        # Search filter
        search = self.request.GET.get('search', '')
        if search:
            queryset = queryset.filter(
                db_models.Q(treatment_code__icontains=search) |
                db_models.Q(name__icontains=search)
            )
        
        # Status filter
        status = self.request.GET.get('status', '')
        if status == 'active':
            queryset = queryset.filter(is_active=True)
        elif status == 'inactive':
            queryset = queryset.filter(is_active=False)
        
        return queryset.order_by('treatment_code')
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['master_type'] = 'nursetreatment'
        context['master_name'] = 'Nurse Treatment Master'
        context['required_fields'] = ['name']
        context['optional_fields'] = ['number_of_nurse_required', 'is_active']
        context['search'] = self.request.GET.get('search', '')
        context['status_filter'] = self.request.GET.get('status', '')
        return context
    
@method_decorator(super_admin_required, name='dispatch')
class NurseTreatmentMasterCreateView(CreateView):
    model = NurseTreatmentMaster
    fields = ['name', 'number_of_nurse_required', 'is_active']
    template_name = 'masters/nursetreatment_form.html'
    success_url = reverse_lazy('nursetreatment_manage')
    
    def form_valid(self, form):
        # Auto-generate unique treatment code by incrementing from highest existing code
        # Get all existing treatment codes that start with NT
        existing_codes = NurseTreatmentMaster.objects.filter(
            treatment_code__startswith='NT'
        ).values_list('treatment_code', flat=True)
        
        # Extract numbers from existing codes and find the highest
        max_number = 0
        for code in existing_codes:
            try:
                # Extract number part after 'NT'
                number_part = int(code.replace('NT', ''))
                if number_part > max_number:
                    max_number = number_part
            except (ValueError, AttributeError):
                continue
        
        # Increment by 1
        next_number = max_number + 1
        
        # Ensure it doesn't exceed 999
        if next_number > 999:
            raise ValueError("Maximum treatment code limit reached (NT999)")
        
        # Generate new code
        treatment_code = f"NT{next_number:03d}"
        
        # Set the auto-generated code
        instance = form.save(commit=False)
        instance.treatment_code = treatment_code
        instance.save()
        _log_master_audit(self.request, 'create', 'NurseTreatmentMaster', str(instance.pk), f'Created NurseTreatmentMaster: {instance.name} ({treatment_code})')
        messages.success(self.request, f'Nurse treatment "{instance.name}" ({treatment_code}) created successfully.')
        return redirect(self.success_url)

@method_decorator(super_admin_required, name='dispatch')
class NurseTreatmentMasterUpdateView(UpdateView):
    model = NurseTreatmentMaster
    fields = ['name', 'number_of_nurse_required', 'is_active']
    template_name = 'masters/nursetreatment_form.html'
    success_url = reverse_lazy('nursetreatment_manage')
    
    def form_valid(self, form):
        _log_master_audit(self.request, 'update', 'NurseTreatmentMaster', str(self.object.pk), f'Updated NurseTreatmentMaster: {self.object.name}')
        messages.success(self.request, 'Nurse treatment updated successfully.')
        return super().form_valid(form)

@login_required
@super_admin_required
def nurse_treatment_delete(request, pk):
    """Delete a nurse treatment - function-based view to handle GET requests properly"""
    treatment = get_object_or_404(NurseTreatmentMaster, pk=pk)
    treatment_name = f"{treatment.treatment_code} - {treatment.name}"
    treatment.delete()
    _log_master_audit(request, 'delete', 'NurseTreatmentMaster', str(pk), f'Deleted NurseTreatmentMaster: {treatment_name}')
    messages.success(request, f'Nurse treatment "{treatment_name}" deleted successfully.')
    return redirect('nursetreatment_manage')

@login_required
@super_admin_required
def nurse_treatment_enable(request, pk):
    """Enable a nurse treatment"""
    treatment = get_object_or_404(NurseTreatmentMaster, pk=pk)
    treatment.is_active = True
    treatment.save()
    _log_master_audit(request, 'update', 'NurseTreatmentMaster', str(pk), f'Enabled NurseTreatmentMaster: {treatment.treatment_code} - {treatment.name}')
    messages.success(request, f'Nurse treatment "{treatment.treatment_code} - {treatment.name}" enabled successfully.')
    return redirect('nursetreatment_manage')

@login_required
@super_admin_required
def nurse_treatment_disable(request, pk):
    """Disable a nurse treatment"""
    treatment = get_object_or_404(NurseTreatmentMaster, pk=pk)
    treatment.is_active = False
    treatment.save()
    _log_master_audit(request, 'update', 'NurseTreatmentMaster', str(pk), f'Disabled NurseTreatmentMaster: {treatment.treatment_code} - {treatment.name}')
    messages.success(request, f'Nurse treatment "{treatment.treatment_code} - {treatment.name}" disabled successfully.')
    return redirect('nursetreatment_manage')

# Nurse Master Management Views
@method_decorator(super_admin_required, name='dispatch')
class NurseMasterManagementView(ListView):
    model = NurseMaster
    template_name = 'masters/nursemaster_manage.html'
    context_object_name = 'nursemaster_list'
    paginate_by = 20

    def get_queryset(self):
        queryset = NurseMaster.objects.all()
        
        # Search filter
        search = self.request.GET.get('search', '')
        if search:
            queryset = queryset.filter(
                db_models.Q(nurse_code__icontains=search) |
                db_models.Q(nurse_name__icontains=search) |
                db_models.Q(designation__icontains=search)
            )
        
        # Status filter
        status = self.request.GET.get('status', '')
        if status == 'active':
            queryset = queryset.filter(is_active=True)
        elif status == 'inactive':
            queryset = queryset.filter(is_active=False)
        
        return queryset.order_by('nurse_code')
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['master_type'] = 'nursemaster'
        context['master_name'] = 'Nurse Master'
        context['required_fields'] = ['nurse_name']
        context['optional_fields'] = ['designation', 'is_active']

        context['search'] = self.request.GET.get('search', '')
        context['status_filter'] = self.request.GET.get('status', '')
        return context
    
    def post(self, request, *args, **kwargs):
        """Handle POST requests for creating new nurses"""
        nurse_name = request.POST.get('nurse_name')
        designation = request.POST.get('designation', '')
        is_active = request.POST.get('is_active') == 'on'
        
        # Create new record
        if nurse_name:
            try:
                # Auto-generate nurse_code (numeric 3-5 digits)
                nurse_code = generate_numeric_code(NurseMaster, 'nurse_code', start_from=100)
                obj = NurseMaster.objects.create(
                    nurse_code=nurse_code,
                    nurse_name=nurse_name,
                    designation=designation,
                    is_active=is_active
                )
                _log_master_audit(request, 'create', 'NurseMaster', str(obj.id), f'Created NurseMaster: {nurse_code} - {nurse_name}')
                messages.success(request, f'Nurse "{nurse_code} - {nurse_name}" created successfully.')
            except Exception as e:
                messages.error(request, f'Error creating nurse: {str(e)}')
        else:
            messages.error(request, 'Nurse name is required.')
        
        return redirect('nursemaster_manage')


@method_decorator(super_admin_required, name='dispatch')
class NurseMasterUpdateView(UpdateView):
    model = NurseMaster
    fields = ['nurse_code', 'nurse_name', 'designation', 'is_active']
    template_name = 'masters/nursemaster_form.html'
    success_url = reverse_lazy('nursemaster_manage')
    
    def get_form(self, form_class=None):
        form = super().get_form(form_class)
        if 'nurse_code' in form.fields:
            form.fields['nurse_code'].disabled = True
            form.fields['nurse_code'].widget.attrs['class'] = 'form-control'
        return form
    
    def form_valid(self, form):
        _log_master_audit(self.request, 'update', 'NurseMaster', str(self.object.pk), f'Updated NurseMaster: {self.object.nurse_code} - {self.object.nurse_name}')
        messages.success(self.request, 'Nurse updated successfully.')
        return super().form_valid(form)


@login_required
@super_admin_required
def nurse_master_delete(request, pk):
    """Delete a nurse - function-based view to handle GET requests properly"""
    nurse = get_object_or_404(NurseMaster, pk=pk)
    nurse_name = f"{nurse.nurse_code} - {nurse.nurse_name}"
    nurse.delete()
    _log_master_audit(request, 'delete', 'NurseMaster', str(pk), f'Deleted NurseMaster: {nurse_name}')
    messages.success(request, f'Nurse "{nurse_name}" deleted successfully.')
    return redirect('nursemaster_manage')

@login_required
@super_admin_required
def nurse_master_enable(request, pk):
    """Enable a nurse"""
    nurse = get_object_or_404(NurseMaster, pk=pk)
    nurse.is_active = True
    nurse.save()
    _log_master_audit(request, 'update', 'NurseMaster', str(pk), f'Enabled NurseMaster: {nurse.nurse_code} - {nurse.nurse_name}')
    messages.success(request, f'Nurse "{nurse.nurse_code} - {nurse.nurse_name}" enabled successfully.')
    return redirect('nursemaster_manage')

@login_required
@super_admin_required
def nurse_master_disable(request, pk):
    """Disable a nurse"""
    nurse = get_object_or_404(NurseMaster, pk=pk)
    nurse.is_active = False
    nurse.save()
    _log_master_audit(request, 'update', 'NurseMaster', str(pk), f'Disabled NurseMaster: {nurse.nurse_code} - {nurse.nurse_name}')
    messages.success(request, f'Nurse "{nurse.nurse_code} - {nurse.nurse_name}" disabled successfully.')
    return redirect('nursemaster_manage')

class PrescriptionTimeManagementView(ListView):
    model = PrescriptionTime
    template_name = 'masters/prestime_manage.html'
    context_object_name = 'prestime_list'
    paginate_by = 20

    def get_queryset(self):
        queryset = PrescriptionTime.objects.all()
        
        # Search filter
        search = self.request.GET.get('search', '')
        if search:
            queryset = queryset.filter(name__icontains=search)
        
        # Status filter
        status = self.request.GET.get('status', '')
        if status == 'active':
            queryset = queryset.filter(is_active=True)
        elif status == 'inactive':
            queryset = queryset.filter(is_active=False)
        return queryset.order_by('order_no', 'name')
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['master_type'] = 'prescriptiontime'
        context['master_name'] = 'Prescription Times'
        context['required_fields'] = ['name']
        context['optional_fields'] = ['is_active']
        context['search'] = self.request.GET.get('search', '')
        context['status_filter'] = self.request.GET.get('status', '')
        return context
    
    def post(self, request, *args, **kwargs):
        """Handle POST requests for creating new prescription times"""
        name = request.POST.get('name')
        order_no = request.POST.get('order_no', 0)
        try:
            order_no = int(order_no)
        except ValueError:
            order_no = 0
        is_active = request.POST.get('is_active') == 'on'
        
        # Create new record
        if name:
            try:
                obj = PrescriptionTime.objects.create(name=name, order_no=order_no, is_active=is_active)
                _log_master_audit(request, 'create', 'PrescriptionTime', str(obj.id), f'Created PrescriptionTime: {name}')
                messages.success(request, f'Prescription time "{name}" created successfully.')
            except Exception as e:
                messages.error(request, f'Error creating prescription time: {str(e)}')
        else:
            messages.error(request, 'Name is required.')
        
        return redirect('prestime_manage')

class PrescriptionTimeUpdateView(UpdateView):
    model = PrescriptionTime
    fields = ['name', 'order_no', 'is_active']
    template_name = 'masters/prestime_form.html'
    success_url = reverse_lazy('prestime_manage')
    
    def form_valid(self, form):
        _log_master_audit(self.request, 'update', 'PrescriptionTime', str(self.object.pk), f'Updated PrescriptionTime: {self.object.name}')
        messages.success(self.request, 'Prescription time updated successfully.')
        return super().form_valid(form)

@login_required
@super_admin_required
def prescription_time_delete(request, pk):
    """Delete a prescription time - function-based view to handle GET requests properly"""
    time = get_object_or_404(PrescriptionTime, pk=pk)
    time_name = time.name
    time.delete()
    _log_master_audit(request, 'delete', 'PrescriptionTime', str(pk), f'Deleted PrescriptionTime: {time_name}')
    messages.success(request, f'Prescription time "{time_name}" deleted successfully.')
    return redirect('prestime_manage')

@login_required
@super_admin_required
def prescription_time_enable(request, pk):
    """Enable a prescription time"""
    time = get_object_or_404(PrescriptionTime, pk=pk)
    time.is_active = True
    time.save()
    _log_master_audit(request, 'update', 'PrescriptionTime', str(pk), f'Enabled PrescriptionTime: {time.name}')
    messages.success(request, f'Prescription time "{time.name}" enabled successfully.')
    return redirect('prestime_manage')

@login_required
@super_admin_required
def prescription_time_disable(request, pk):
    """Disable a prescription time"""
    time = get_object_or_404(PrescriptionTime, pk=pk)
    time.is_active = False
    time.save()
    _log_master_audit(request, 'update', 'PrescriptionTime', str(pk), f'Disabled PrescriptionTime: {time.name}')
    messages.success(request, f'Prescription time "{time.name}" disabled successfully.')
    return redirect('prestime_manage')

class PrescriptionScopeManagementView(ListView):
    model = PrescriptionScope
    template_name = 'masters/prescope_manage.html'
    context_object_name = 'prescope_list'
    paginate_by = 20

    def get_queryset(self):
        queryset = PrescriptionScope.objects.all()
        
        # Search filter
        search = self.request.GET.get('search', '')
        if search:
            queryset = queryset.filter(name__icontains=search)
        
        # Status filter
        status = self.request.GET.get('status', '')
        if status == 'active':
            queryset = queryset.filter(is_active=True)
        elif status == 'inactive':
            queryset = queryset.filter(is_active=False)
        
        return queryset.order_by('name')
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['master_type'] = 'prescriptionscope'
        context['master_name'] = 'Prescription Scopes'
        context['required_fields'] = ['name']
        context['optional_fields'] = ['is_active']
        context['search'] = self.request.GET.get('search', '')
        context['status_filter'] = self.request.GET.get('status', '')
        return context
    
    def post(self, request, *args, **kwargs):
        """Handle POST requests for creating new prescription scopes"""
        name = request.POST.get('name')
        is_active = request.POST.get('is_active') == 'on'
        
        # Create new record
        if name:
            try:
                obj = PrescriptionScope.objects.create(name=name, is_active=is_active)
                _log_master_audit(request, 'create', 'PrescriptionScope', str(obj.id), f'Created PrescriptionScope: {name}')
                messages.success(request, f'Prescription scope "{name}" created successfully.')
            except Exception as e:
                messages.error(request, f'Error creating prescription scope: {str(e)}')
        else:
            messages.error(request, 'Name is required.')
        
        return redirect('prescope_manage')

class PrescriptionScopeUpdateView(UpdateView):
    model = PrescriptionScope
    fields = ['name', 'is_active']
    template_name = 'masters/prescope_form.html'
    success_url = reverse_lazy('prescope_manage')
    
    def form_valid(self, form):
        _log_master_audit(self.request, 'update', 'PrescriptionScope', str(self.object.pk), f'Updated PrescriptionScope: {self.object.name}')
        messages.success(self.request, 'Prescription scope updated successfully.')
        return super().form_valid(form)

@login_required
@super_admin_required
def prescription_scope_delete(request, pk):
    """Delete a prescription scope - function-based view to handle GET requests properly"""
    scope = get_object_or_404(PrescriptionScope, pk=pk)
    scope_name = scope.name
    scope.delete()
    _log_master_audit(request, 'delete', 'PrescriptionScope', str(pk), f'Deleted PrescriptionScope: {scope_name}')
    messages.success(request, f'Prescription scope "{scope_name}" deleted successfully.')
    return redirect('prescope_manage')

@login_required
@super_admin_required
def prescription_scope_enable(request, pk):
    """Enable a prescription scope"""
    scope = get_object_or_404(PrescriptionScope, pk=pk)
    scope.is_active = True
    scope.save()
    _log_master_audit(request, 'update', 'PrescriptionScope', str(pk), f'Enabled PrescriptionScope: {scope.name}')
    messages.success(request, f'Prescription scope "{scope.name}" enabled successfully.')
    return redirect('prescope_manage')

@login_required
@super_admin_required
def prescription_scope_disable(request, pk):
    """Disable a prescription scope"""
    scope = get_object_or_404(PrescriptionScope, pk=pk)
    scope.is_active = False
    scope.save()
    _log_master_audit(request, 'update', 'PrescriptionScope', str(pk), f'Disabled PrescriptionScope: {scope.name}')
    messages.success(request, f'Prescription scope "{scope.name}" disabled successfully.')
    return redirect('prescope_manage')

# UnitNo views
class UnitNoListView(ListView):
    model = UnitNo
    template_name = 'masters/unitno_master.html'
    context_object_name = 'unitno_list'
    paginate_by = 20

    def get_queryset(self):
        queryset = UnitNo.objects.all()
        
        # Search filter
        search = self.request.GET.get('search', '')
        if search:
            queryset = queryset.filter(unit_name__icontains=search)
        
        # Status filter
        status = self.request.GET.get('status', '')
        if status == 'active':
            queryset = queryset.filter(is_active=True)
        elif status == 'inactive':
            queryset = queryset.filter(is_active=False)
        
        return queryset.order_by('unit_name')
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['master_type'] = 'unitno'
        context['master_name'] = 'Unit Numbers'
        context['required_fields'] = ['unit_name']
        context['optional_fields'] = ['is_active']
        context['search'] = self.request.GET.get('search', '')
        context['status_filter'] = self.request.GET.get('status', '')
        return context
    
    def post(self, request, *args, **kwargs):
        """Handle POST requests for creating new units"""
        unit_name = request.POST.get('unit_name')
        is_active = request.POST.get('is_active') == 'on'
        
        if unit_name:
            try:
                obj = UnitNo.objects.create(unit_name=unit_name, is_active=is_active)
                _log_master_audit(request, 'create', 'UnitNo', str(obj.id), f'Created UnitNo: {unit_name}')
                messages.success(request, f'Unit "{unit_name}" created successfully.')
            except Exception as e:
                messages.error(request, f'Error creating unit: {str(e)}')
        else:
            messages.error(request, 'Unit name is required.')
        
        return redirect('unit_list')

class UnitNoCreateView(CreateView):
    model = UnitNo
    fields = ['unit_name', 'is_active']
    template_name = 'masters/unitno_form.html'
    success_url = reverse_lazy('unit_list')
    
    def form_valid(self, form):
        response = super().form_valid(form)
        _log_master_audit(self.request, 'create', 'UnitNo', str(self.object.pk), f'Created UnitNo: {self.object.unit_name}')
        messages.success(self.request, 'Unit created successfully.')
        return response

class UnitNoUpdateView(UpdateView):
    model = UnitNo
    fields = ['unit_name', 'is_active']
    template_name = 'masters/unitno_form.html'
    success_url = reverse_lazy('unit_list')
    
    def form_valid(self, form):
        _log_master_audit(self.request, 'update', 'UnitNo', str(self.object.pk), f'Updated UnitNo: {self.object.unit_name}')
        messages.success(self.request, 'Unit updated successfully.')
        return super().form_valid(form)

@login_required
@super_admin_required
def unit_no_delete(request, pk):
    """Delete a unit - function-based view to handle GET requests properly"""
    unit = get_object_or_404(UnitNo, pk=pk)
    unit_name = unit.unit_name
    unit.delete()
    _log_master_audit(request, 'delete', 'UnitNo', str(pk), f'Deleted UnitNo: {unit_name}')
    messages.success(request, f'Unit "{unit_name}" deleted successfully.')
    return redirect('unit_list')

@login_required
@super_admin_required
def unit_no_enable(request, pk):
    unit = get_object_or_404(UnitNo, pk=pk)
    unit.is_active = True
    unit.save()
    _log_master_audit(request, 'update', 'UnitNo', str(pk), f'Enabled UnitNo: {unit.unit_name}')
    messages.success(request, f'Unit {unit.unit_name} enabled successfully.')
    return redirect('unit_list')

@login_required
@super_admin_required
def unit_no_disable(request, pk):
    unit = get_object_or_404(UnitNo, pk=pk)
    unit.is_active = False
    unit.save()
    _log_master_audit(request, 'update', 'UnitNo', str(pk), f'Disabled UnitNo: {unit.unit_name}')
    messages.success(request, f'Unit {unit.unit_name} disabled successfully.')
    return redirect('unit_list')

# Master Management Dashboard
@login_required
@super_admin_required
def master_management_dashboard(request):
    """Comprehensive master management dashboard for Super Admin"""
    context = {
        'master_tables': [
            {
                'name': 'Therapy Master',
                'model': 'TherapyMaster',
                'url': 'therapy_master_manage',
                'description': 'Manage therapy types and details',
                'count': TherapyMaster.objects.count(),
                'active_count': TherapyMaster.objects.filter(is_active=True).count(),
            },
            {
                'name': 'Therapy Group Master',
                'model': 'TherapyGroup',
                'url': 'therapy_group_manage',
                'description': 'Manage therapy groups/categories',
                'count': TherapyGroup.objects.count(),
                'active_count': TherapyGroup.objects.filter(is_active=True).count(),
            },
            {
                'name': 'Item Master',
                'model': 'ItemMaster',
                'url': 'item_master_manage',
                'description': 'Manage items used in therapies',
                'count': ItemMaster.objects.count(),
                'active_count': ItemMaster.objects.filter(is_active=True).count(),
            },
            {
                'name': 'Item Group Master',
                'model': 'ItemGroup',
                'url': 'item_group_manage',
                'description': 'Manage item groups/categories',
                'count': ItemGroup.objects.count(),
                'active_count': ItemGroup.objects.filter(is_active=True).count(),
            },
            {
                'name': 'Slot Master',
                'model': 'SlotMaster',
                'url': 'slot_master_manage',
                'description': 'Manage therapy time slots',
                'count': SlotMaster.objects.count(),
                'active_count': SlotMaster.objects.filter(is_active=True).count(),
            },
            {
                'name': 'Therapist Master',
                'model': 'TherapistMaster',
                'url': 'therapist_master_manage',
                'description': 'Manage therapists and their details',
                'count': TherapistMaster.objects.count(),
                'active_count': TherapistMaster.objects.filter(is_active=True).count(),
            },
            {
                'name': 'Procurement Sources',
                'model': 'Procured',
                'url': 'procured_manage',
                'description': 'Manage procurement sources',
                'count': Procured.objects.count(),
                'active_count': Procured.objects.filter(is_active=True).count(),
            },
            {
                'name': 'Oil Classes',
                'model': 'OilClass',
                'url': 'oilclass_manage',
                'description': 'Manage oil categories',
                'count': OilClass.objects.count(),
                'active_count': OilClass.objects.filter(is_active=True).count(),
            },
            {
                'name': 'Usage Types',
                'model': 'UsageType',
                'url': 'usagetype_manage',
                'description': 'Manage usage types',
                'count': UsageType.objects.count(),
                'active_count': UsageType.objects.filter(is_active=True).count(),
            },
            {
                'name': 'Prescription Times',
                'model': 'PrescriptionTime',
                'url': 'prestime_manage',
                'description': 'Manage prescription time slots',
                'count': PrescriptionTime.objects.count(),
                'active_count': PrescriptionTime.objects.filter(is_active=True).count(),
            },
            {
                'name': 'Prescription Scopes',
                'model': 'PrescriptionScope',
                'url': 'prescope_manage',
                'description': 'Manage prescription scopes',
                'count': PrescriptionScope.objects.count(),
                'active_count': PrescriptionScope.objects.filter(is_active=True).count(),
            },
            {
                'name': 'Unit Numbers',
                'model': 'UnitNo',
                'url': 'unit_list',
                'description': 'Manage unit measurements',
                'count': UnitNo.objects.count(),
                'active_count': UnitNo.objects.filter(is_active=True).count(),
            },
            {
                'name': 'Ingredient Master',
                'model': 'IngredientMaster',
                'url': 'ingredient_master_manage',
                'description': 'Manage ingredients used in prescriptions',
                'count': IngredientMaster.objects.count(),
                'active_count': IngredientMaster.objects.filter(is_active=True).count(),
            },
            {
                'name': 'Treatment Room Master',
                'model': 'TreatmentRoom',
                'url': 'treatment_room_master_manage',
                'description': 'Manage treatment rooms for therapy allocation',
                'count': TreatmentRoom.objects.count(),
                'active_count': TreatmentRoom.objects.filter(is_active=True).count(),
            },
        ]
    }
    return render(request, 'masters/master_management_dashboard.html', context)

# Therapy Master CRUD Views
class TherapyMasterManagementView(ListView):
    model = TherapyMaster
    template_name = 'masters/therapy_master_manage.html'
    context_object_name = 'therapy_list'
    paginate_by = 20

    def get_queryset(self):
        queryset = TherapyMaster.objects.all()
        
        # Search filter (server-side)
        search = self.request.GET.get('search', '')
        if search:
            queryset = queryset.filter(
                db_models.Q(therapy_code__icontains=search) |
                db_models.Q(therapy_name__icontains=search) |
                db_models.Q(therapy_group__name__icontains=search)
            )
        
        # Status filter (server-side)
        status = self.request.GET.get('status', '')
        if status == 'active':
            queryset = queryset.filter(is_active=True)
        elif status == 'inactive':
            queryset = queryset.filter(is_active=False)
        
        # Category filter (server-side)
        category = self.request.GET.get('category', '')
        if category:
            queryset = queryset.filter(therapy_category=category)
        
        return queryset.order_by('therapy_name')
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['search'] = self.request.GET.get('search', '')
        context['status_filter'] = self.request.GET.get('status', '')
        context['category_filter'] = self.request.GET.get('category', '')
        # Bulk upload context
        context['master_type'] = 'therapy'
        context['master_name'] = 'Therapy Master'
        context['required_fields'] = ['therapy_code', 'therapy_name', 'therapy_category', 'therapy_duration',
                                     'bathing_required', 'treatment_room_required', 'cost_of_therapy', 'patient_type',
                                     'no_of_super_senior_therapist', 'no_of_senior_therapist',
                                     'no_of_junior_therapist', 'no_of_trainee_therapist', 'is_active']
        context['optional_fields'] = ['therapy_group']
        return context

class TherapyMasterCreateView(CreateView):
    model = TherapyMaster
    fields = ['therapy_name', 'therapy_group', 'therapy_category', 'therapy_duration', 'bathing_required', 'treatment_room_required', 'cost_of_therapy', 'patient_type', 'no_of_super_senior_therapist', 'no_of_senior_therapist', 'no_of_junior_therapist', 'no_of_trainee_therapist', 'is_active']
    template_name = 'masters/therapy_master_form.html'
    success_url = reverse_lazy('therapy_master_manage')
    
    def get_form(self, form_class=None):
        form = super().get_form(form_class)
        # Make therapy_code read-only/hidden since it's auto-generated
        if 'therapy_code' in form.fields:
            form.fields['therapy_code'].widget.attrs['readonly'] = True
        # Filter therapy_group to show only active groups - keep optional (only non-required field)
        if 'therapy_group' in form.fields:
            form.fields['therapy_group'].queryset = TherapyGroup.objects.filter(is_active=True).order_by('name')
            form.fields['therapy_group'].empty_label = '-- Select Group --'
            form.fields['therapy_group'].widget.attrs['class'] = 'form-control'
            form.fields['therapy_group'].required = False
        # Make all other fields required (except therapy_group)
        for field_name in ['therapy_name', 'therapy_category', 'therapy_duration', 'bathing_required',
                           'treatment_room_required', 'cost_of_therapy', 'patient_type',
                           'no_of_super_senior_therapist', 'no_of_senior_therapist',
                           'no_of_junior_therapist', 'no_of_trainee_therapist']:
            if field_name in form.fields:
                form.fields[field_name].required = True
        # Update therapy_category to use dropdown with Medium, Hard, Easy
        if 'therapy_category' in form.fields:
            form.fields['therapy_category'].widget = forms.Select(choices=[
                ('', 'Select Category'),
                ('Medium', 'Medium'),
                ('Hard', 'Hard'),
                ('Easy', 'Easy'),
            ], attrs={'class': 'form-select'})
        return form
    
    def form_valid(self, form):
        # Auto-generate unique therapy code by incrementing from highest existing code
        # Get all existing therapy codes that start with THEB
        existing_codes = TherapyMaster.objects.filter(
            therapy_code__startswith='THEB'
        ).values_list('therapy_code', flat=True)
        
        # Extract numbers from existing codes and find the highest
        max_number = 0
        for code in existing_codes:
            try:
                # Extract number part after 'THEB'
                number_part = int(code.replace('THEB', ''))
                if number_part > max_number:
                    max_number = number_part
            except (ValueError, AttributeError):
                continue
        
        # Increment by 1
        next_number = max_number + 1
        
        # Ensure it doesn't exceed 999
        if next_number > 999:
            raise ValueError("Maximum therapy code limit reached (THEB999)")
        
        # Generate new code
        therapy_code = f"THEB{next_number:03d}"
        
        # Calculate total number of therapists
        instance = form.save(commit=False)
        instance.therapy_code = therapy_code
        instance.total_no_of_therapist = (
            instance.no_of_super_senior_therapist +
            instance.no_of_senior_therapist +
            instance.no_of_junior_therapist +
            instance.no_of_trainee_therapist
        )
        instance.save()
        _log_master_audit(self.request, 'create', 'TherapyMaster', str(instance.pk), f'Created TherapyMaster: {instance.therapy_name} ({therapy_code})')
        messages.success(self.request, f'Therapy "{instance.therapy_name}" ({therapy_code}) created successfully.')
        return redirect(self.success_url)

class TherapyMasterUpdateView(UpdateView):
    model = TherapyMaster
    fields = ['therapy_name', 'therapy_group', 'therapy_category', 'therapy_duration', 'bathing_required', 'treatment_room_required', 'cost_of_therapy', 'patient_type', 'no_of_super_senior_therapist', 'no_of_senior_therapist', 'no_of_junior_therapist', 'no_of_trainee_therapist', 'is_active']
    template_name = 'masters/therapy_master_form.html'
    success_url = reverse_lazy('therapy_master_manage')
    
    def get_form(self, form_class=None):
        form = super().get_form(form_class)
        # Make therapy_code read-only since it shouldn't be changed
        if 'therapy_code' in form.fields:
            form.fields['therapy_code'].widget.attrs['readonly'] = True
        # Filter therapy_group to show only active groups (and current group if inactive) - keep optional
        if 'therapy_group' in form.fields:
            current_group = self.object.therapy_group
            if current_group:
                # Include current group even if inactive
                form.fields['therapy_group'].queryset = TherapyGroup.objects.filter(
                    db_models.Q(is_active=True) | db_models.Q(pk=current_group.pk)
                ).order_by('name')
            else:
                form.fields['therapy_group'].queryset = TherapyGroup.objects.filter(is_active=True).order_by('name')
            form.fields['therapy_group'].empty_label = '-- Select Group --'
            form.fields['therapy_group'].widget.attrs['class'] = 'form-control'
            form.fields['therapy_group'].required = False
        # Make all other fields required (except therapy_group)
        for field_name in ['therapy_name', 'therapy_category', 'therapy_duration', 'bathing_required',
                           'treatment_room_required', 'cost_of_therapy', 'patient_type',
                           'no_of_super_senior_therapist', 'no_of_senior_therapist',
                           'no_of_junior_therapist', 'no_of_trainee_therapist']:
            if field_name in form.fields:
                form.fields[field_name].required = True
        # Update therapy_category to use dropdown with Medium, Hard, Easy
        if 'therapy_category' in form.fields:
            form.fields['therapy_category'].widget = forms.Select(choices=[
                ('', 'Select Category'),
                ('Medium', 'Medium'),
                ('Hard', 'Hard'),
                ('Easy', 'Easy'),
            ], attrs={'class': 'form-select'})
        return form
    
    def form_valid(self, form):
        # Auto-calculate total number of therapists
        instance = form.save(commit=False)
        instance.total_no_of_therapist = (
            instance.no_of_super_senior_therapist +
            instance.no_of_senior_therapist +
            instance.no_of_junior_therapist +
            instance.no_of_trainee_therapist
        )
        instance.save()
        _log_master_audit(self.request, 'update', 'TherapyMaster', str(self.object.pk), f'Updated TherapyMaster: {instance.therapy_name}')
        messages.success(self.request, f'Therapy "{instance.therapy_name}" updated successfully.')
        return redirect(self.success_url)

@login_required
@super_admin_required
def therapy_master_delete(request, pk):
    """Delete a therapy - function-based view to handle GET requests properly"""
    therapy = get_object_or_404(TherapyMaster, pk=pk)
    therapy_name = therapy.therapy_name
    therapy.delete()
    _log_master_audit(request, 'delete', 'TherapyMaster', str(pk), f'Deleted TherapyMaster: {therapy_name}')
    messages.success(request, f'Therapy "{therapy_name}" deleted successfully.')
    return redirect('therapy_master_manage')


@login_required
@super_admin_required
def therapy_master_enable(request, pk):
    """Enable a therapy master record."""
    therapy = get_object_or_404(TherapyMaster, pk=pk)
    therapy.is_active = True
    therapy.save()
    _log_master_audit(request, 'update', 'TherapyMaster', str(pk), f'Enabled TherapyMaster: {therapy.therapy_name}')
    messages.success(request, f'Therapy "{therapy.therapy_name}" enabled successfully.')
    return redirect('therapy_master_manage')


@login_required
@super_admin_required
def therapy_master_disable(request, pk):
    """Disable a therapy master record."""
    therapy = get_object_or_404(TherapyMaster, pk=pk)
    therapy.is_active = False
    therapy.save()
    _log_master_audit(request, 'update', 'TherapyMaster', str(pk), f'Disabled TherapyMaster: {therapy.therapy_name}')
    messages.success(request, f'Therapy "{therapy.therapy_name}" disabled successfully.')
    return redirect('therapy_master_manage')

# Item Master CRUD Views
class ItemMasterManagementView(ListView):
    model = ItemMaster
    template_name = 'masters/item_master_manage.html'
    context_object_name = 'item_list'
    paginate_by = 20

    def get_queryset(self):
        queryset = ItemMaster.objects.all()
        
        # Search filter
        search = self.request.GET.get('search', '')
        if search:
            queryset = queryset.filter(
                db_models.Q(item_code__icontains=search) |
                db_models.Q(item_name__icontains=search) |
                db_models.Q(item_group__name__icontains=search) |
                db_models.Q(procured_from__name__icontains=search)
            )
        
        # Status filter
        status = self.request.GET.get('status', '')
        if status == 'active':
            queryset = queryset.filter(is_active=True)
        elif status == 'inactive':
            queryset = queryset.filter(is_active=False)
        
        # Group filter
        group = self.request.GET.get('group', '')
        if group:
            queryset = queryset.filter(item_group_id=group)
        
        return queryset.order_by('item_name')
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['master_type'] = 'item'
        context['master_name'] = 'Item Master'
        context['required_fields'] = ['item_name']
        context['optional_fields'] = ['item_group', 'umo', 'min_quantity', 'max_quantity', 'procured_from', 'is_active']
        context['search'] = self.request.GET.get('search', '')
        context['status_filter'] = self.request.GET.get('status', '')
        context['group_filter'] = self.request.GET.get('group', '')
        context['item_groups'] = ItemGroup.objects.filter(is_active=True).order_by('name')
        return context


class ItemMasterForm(forms.ModelForm):
    """Custom form for ItemMaster with filtered procured_from dropdown"""
    class Meta:
        model = ItemMaster
        fields = ['item_name', 'item_group', 'umo', 'min_quantity', 'max_quantity', 'procured_from', 'is_active']
        widgets = {
            'item_name': forms.TextInput(attrs={'class': 'form-control'}),
            'item_group': forms.Select(attrs={'class': 'form-control'}),
            'umo': forms.Select(attrs={'class': 'form-control'}),
            'min_quantity': forms.NumberInput(attrs={'class': 'form-control', 'min': 0, 'placeholder': 'Min'}),
            'max_quantity': forms.NumberInput(attrs={'class': 'form-control', 'min': 0, 'placeholder': 'Max'}),
            'procured_from': forms.Select(attrs={'class': 'form-control'}),
            'is_active': forms.CheckboxInput(attrs={'class': 'form-check-input'}),
        }
    
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        # Filter procured_from to show only active sources
        self.fields['procured_from'].queryset = Procured.objects.filter(is_active=True).order_by('name')
        self.fields['procured_from'].empty_label = '-- Select Source --'
        # Filter item_group to show only active groups
        self.fields['item_group'].queryset = ItemGroup.objects.filter(is_active=True).order_by('name')
        self.fields['item_group'].empty_label = '-- Select Group --'
        # Filter umo to show only active units
        current_umo = self.instance.umo if self.instance and self.instance.pk else None
        if current_umo:
            # Include current unit even if inactive
            self.fields['umo'].queryset = UnitNo.objects.filter(
                db_models.Q(is_active=True) | db_models.Q(pk=current_umo.pk)
            ).order_by('unit_name')
        else:
            self.fields['umo'].queryset = UnitNo.objects.filter(is_active=True).order_by('unit_name')
        self.fields['umo'].empty_label = '-- Select UOM --'
        # Show item_code as read-only if editing
        if self.instance and self.instance.pk:
            self.fields['item_code'] = forms.CharField(
                widget=forms.TextInput(attrs={'class': 'form-control', 'readonly': True}),
                required=False,
                initial=self.instance.item_code
            )

class ItemMasterCreateView(CreateView):
    model = ItemMaster
    form_class = ItemMasterForm
    template_name = 'masters/item_master_form.html'
    success_url = reverse_lazy('item_master_manage')
    
    def form_valid(self, form):
        # Auto-generate item_code (numeric 4-5 digits)
        instance = form.save(commit=False)
        instance.item_code = generate_numeric_code(ItemMaster, 'item_code')
        instance.save()
        _log_master_audit(self.request, 'create', 'ItemMaster', str(instance.pk), f'Created ItemMaster: {instance.item_name} ({instance.item_code})')
        messages.success(self.request, f'Item "{instance.item_name}" ({instance.item_code}) created successfully.')
        return redirect(self.success_url)


class ItemMasterUpdateView(UpdateView):
    model = ItemMaster
    form_class = ItemMasterForm
    template_name = 'masters/item_master_form.html'
    success_url = reverse_lazy('item_master_manage')
    
    def form_valid(self, form):
        _log_master_audit(self.request, 'update', 'ItemMaster', str(self.object.pk), f'Updated ItemMaster: {self.object.item_name}')
        messages.success(self.request, 'Item updated successfully.')
        return super().form_valid(form)

@login_required
@super_admin_required
def item_master_delete(request, pk):
    """Delete an item - function-based view to handle GET requests properly"""
    item = get_object_or_404(ItemMaster, pk=pk)
    item_name = item.item_name
    item_code = item.item_code
    item.delete()
    _log_master_audit(request, 'delete', 'ItemMaster', str(pk), f'Deleted ItemMaster: {item_name} ({item_code})')
    messages.success(request, f'Item "{item_name}" ({item_code}) deleted successfully.')
    return redirect('item_master_manage')


@login_required
@super_admin_required
def item_master_enable(request, pk):
    """Enable an item master record."""
    item = get_object_or_404(ItemMaster, pk=pk)
    item.is_active = True
    item.save()
    _log_master_audit(request, 'update', 'ItemMaster', str(pk), f'Enabled ItemMaster: {item.item_name}')
    messages.success(request, f'Item "{item.item_name}" enabled successfully.')
    return redirect('item_master_manage')


@login_required
@super_admin_required
def item_master_disable(request, pk):
    """Disable an item master record."""
    item = get_object_or_404(ItemMaster, pk=pk)
    item.is_active = False
    item.save()
    _log_master_audit(request, 'update', 'ItemMaster', str(pk), f'Disabled ItemMaster: {item.item_name}')
    messages.success(request, f'Item "{item.item_name}" disabled successfully.')
    return redirect('item_master_manage')

# Item Group Master CRUD Views
class ItemGroupManagementView(ListView):
    model = ItemGroup
    template_name = 'masters/item_group_manage.html'
    context_object_name = 'item_group_list'
    paginate_by = 20

    def get_queryset(self):
        queryset = ItemGroup.objects.all()
        
        # Search filter
        search = self.request.GET.get('search', '').strip()
        if search:
            queryset = queryset.filter(
                db_models.Q(name__icontains=search)
            )
        
        # Status filter
        status = self.request.GET.get('status', '').strip().lower()
        if status == 'active':
            queryset = queryset.filter(is_active=True)
        elif status == 'inactive':
            queryset = queryset.filter(is_active=False)
        
        return queryset.order_by('name')
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['master_type'] = 'item_group'
        context['master_name'] = 'Item Group Master'
        context['required_fields'] = ['name']
        context['optional_fields'] = ['is_active']
        context['search'] = self.request.GET.get('search', '')
        context['status_filter'] = self.request.GET.get('status', '')
        return context
    
    def post(self, request, *args, **kwargs):
        """Handle POST requests for creating new item groups"""
        name = request.POST.get('name')
        is_active = request.POST.get('is_active') == 'on'
        
        # Create new record
        if name:
            try:
                ItemGroup.objects.create(
                    name=name,
                    is_active=is_active
                )
                messages.success(request, f'Item group "{name}" created successfully.')
            except Exception as e:
                messages.error(request, f'Error creating item group: {str(e)}')
        else:
            messages.error(request, 'Name is required.')
        
        return redirect('item_group_manage')

class ItemGroupCreateView(CreateView):
    model = ItemGroup
    fields = ['name', 'is_active']
    template_name = 'masters/item_group_form.html'
    success_url = reverse_lazy('item_group_manage')
    
    def form_valid(self, form):
        response = super().form_valid(form)
        _log_master_audit(self.request, 'create', 'ItemGroup', str(self.object.pk), f'Created ItemGroup: {self.object.name}')
        messages.success(self.request, 'Item group created successfully.')
        return response

class ItemGroupUpdateView(UpdateView):
    model = ItemGroup
    fields = ['name', 'is_active']
    template_name = 'masters/item_group_form.html'
    success_url = reverse_lazy('item_group_manage')
    
    def form_valid(self, form):
        _log_master_audit(self.request, 'update', 'ItemGroup', str(self.object.pk), f'Updated ItemGroup: {self.object.name}')
        messages.success(self.request, 'Item group updated successfully.')
        return super().form_valid(form)

@login_required
@super_admin_required
def item_group_delete(request, pk):
    """Delete an item group - function-based view to handle GET requests properly"""
    try:
        item_group = get_object_or_404(ItemGroup, pk=pk)
        item_group_name = item_group.name
        
        # Check if any items are using this group
        items_using_group = ItemMaster.objects.filter(item_group=item_group).count()
        if items_using_group > 0:
            messages.error(request, f'Cannot delete item group "{item_group_name}" because {items_using_group} item(s) are using it. Please update those items first.')
            return redirect('item_group_manage')
        
        item_group.delete()
        _log_master_audit(request, 'delete', 'ItemGroup', str(pk), f'Deleted ItemGroup: {item_group_name}')
        messages.success(request, f'Item group "{item_group_name}" deleted successfully.')
    except Exception as e:
        messages.error(request, f'Error deleting item group: {str(e)}')
    
    return redirect('item_group_manage')

@login_required
@super_admin_required
def item_group_enable(request, pk):
    """Enable an item group"""
    item_group = get_object_or_404(ItemGroup, pk=pk)
    item_group.is_active = True
    item_group.save()
    _log_master_audit(request, 'update', 'ItemGroup', str(pk), f'Enabled ItemGroup: {item_group.name}')
    messages.success(request, f'Item group "{item_group.name}" enabled successfully.')
    return redirect('item_group_manage')

@login_required
@super_admin_required
def item_group_disable(request, pk):
    """Disable an item group"""
    item_group = get_object_or_404(ItemGroup, pk=pk)
    item_group.is_active = False
    item_group.save()
    _log_master_audit(request, 'update', 'ItemGroup', str(pk), f'Disabled ItemGroup: {item_group.name}')
    messages.success(request, f'Item group "{item_group.name}" disabled successfully.')
    return redirect('item_group_manage')

# Therapy Group Master CRUD Views
class TherapyGroupManagementView(ListView):
    model = TherapyGroup
    template_name = 'masters/therapy_group_manage.html'
    context_object_name = 'therapy_group_list'
    paginate_by = 20

    def get_queryset(self):
        queryset = TherapyGroup.objects.all()
        
        # Search filter
        search = self.request.GET.get('search', '').strip()
        if search:
            queryset = queryset.filter(
                db_models.Q(name__icontains=search)
            )
        
        # Status filter
        status = self.request.GET.get('status', '').strip().lower()
        if status == 'active':
            queryset = queryset.filter(is_active=True)
        elif status == 'inactive':
            queryset = queryset.filter(is_active=False)
        
        return queryset.order_by('name')
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['master_type'] = 'therapy_group'
        context['master_name'] = 'Therapy Group Master'
        context['required_fields'] = ['name']
        context['optional_fields'] = ['is_active']
        context['search'] = self.request.GET.get('search', '')
        context['status_filter'] = self.request.GET.get('status', '')
        return context
    
    def post(self, request, *args, **kwargs):
        """Handle POST requests for creating new therapy groups"""
        name = request.POST.get('name')
        is_active = request.POST.get('is_active') == 'on'
        
        # Create new record
        if name:
            try:
                obj = TherapyGroup.objects.create(name=name, is_active=is_active)
                _log_master_audit(request, 'create', 'TherapyGroup', str(obj.id), f'Created TherapyGroup: {name}')
                messages.success(request, f'Therapy group "{name}" created successfully.')
            except Exception as e:
                messages.error(request, f'Error creating therapy group: {str(e)}')
        else:
            messages.error(request, 'Name is required.')
        
        return redirect('therapy_group_manage')

class TherapyGroupCreateView(CreateView):
    model = TherapyGroup
    fields = ['name', 'is_active']
    template_name = 'masters/therapy_group_form.html'
    success_url = reverse_lazy('therapy_group_manage')
    
    def form_valid(self, form):
        response = super().form_valid(form)
        _log_master_audit(self.request, 'create', 'TherapyGroup', str(self.object.pk), f'Created TherapyGroup: {self.object.name}')
        messages.success(self.request, 'Therapy group created successfully.')
        return response

class TherapyGroupUpdateView(UpdateView):
    model = TherapyGroup
    fields = ['name', 'is_active']
    template_name = 'masters/therapy_group_form.html'
    success_url = reverse_lazy('therapy_group_manage')
    
    def form_valid(self, form):
        _log_master_audit(self.request, 'update', 'TherapyGroup', str(self.object.pk), f'Updated TherapyGroup: {self.object.name}')
        messages.success(self.request, 'Therapy group updated successfully.')
        return super().form_valid(form)

@login_required
@super_admin_required
def therapy_group_delete(request, pk):
    """Delete a therapy group - function-based view to handle GET requests properly"""
    try:
        therapy_group = get_object_or_404(TherapyGroup, pk=pk)
        therapy_group_name = therapy_group.name
        
        # Check if any therapies are using this group
        therapies_using_group = TherapyMaster.objects.filter(therapy_group=therapy_group).count()
        if therapies_using_group > 0:
            messages.error(request, f'Cannot delete therapy group "{therapy_group_name}" because {therapies_using_group} therapy/therapies are using it. Please update those therapies first.')
            return redirect('therapy_group_manage')
        
        therapy_group.delete()
        _log_master_audit(request, 'delete', 'TherapyGroup', str(pk), f'Deleted TherapyGroup: {therapy_group_name}')
        messages.success(request, f'Therapy group "{therapy_group_name}" deleted successfully.')
    except Exception as e:
        messages.error(request, f'Error deleting therapy group: {str(e)}')
    
    return redirect('therapy_group_manage')

@login_required
@super_admin_required
def therapy_group_enable(request, pk):
    """Enable a therapy group"""
    therapy_group = get_object_or_404(TherapyGroup, pk=pk)
    therapy_group.is_active = True
    therapy_group.save()
    _log_master_audit(request, 'update', 'TherapyGroup', str(pk), f'Enabled TherapyGroup: {therapy_group.name}')
    messages.success(request, f'Therapy group "{therapy_group.name}" enabled successfully.')
    return redirect('therapy_group_manage')

@login_required
@super_admin_required
def therapy_group_disable(request, pk):
    """Disable a therapy group"""
    therapy_group = get_object_or_404(TherapyGroup, pk=pk)
    therapy_group.is_active = False
    therapy_group.save()
    _log_master_audit(request, 'update', 'TherapyGroup', str(pk), f'Disabled TherapyGroup: {therapy_group.name}')
    messages.success(request, f'Therapy group "{therapy_group.name}" disabled successfully.')
    return redirect('therapy_group_manage')

# Slot Master CRUD Views
class SlotMasterManagementView(ListView):
    model = SlotMaster
    template_name = 'masters/slot_master_manage.html'
    context_object_name = 'slot_list'
    paginate_by = 20

    def get_queryset(self):
        queryset = SlotMaster.objects.all()
        
        # Search filter
        search = self.request.GET.get('search', '')
        if search:
            queryset = queryset.filter(
                db_models.Q(slot_number__icontains=search) |
                db_models.Q(slot_name__icontains=search)
            )
        
        # Status filter
        status = self.request.GET.get('status', '')
        if status == 'active':
            queryset = queryset.filter(is_active=True)
        elif status == 'inactive':
            queryset = queryset.filter(is_active=False)
        
        return queryset.order_by('slot_number')
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['master_type'] = 'slot'
        context['master_name'] = 'Slot Master'
        context['required_fields'] = ['slot_number', 'slot_time']
        context['optional_fields'] = ['slot_name', 'is_active']
        context['search'] = self.request.GET.get('search', '')
        context['status_filter'] = self.request.GET.get('status', '')
        return context

class SlotMasterCreateView(CreateView):
    model = SlotMaster
    fields = ['slot_name', 'slot_time', 'is_active']
    template_name = 'masters/slot_master_form.html'
    success_url = reverse_lazy('slot_master_manage')
    
    def form_valid(self, form):
        # Auto-generate slot_number
        with transaction.atomic():
            max_val = SlotMaster.objects.aggregate(Max('slot_number'))['slot_number__max']
            next_val = (max_val or 0) + 1
            instance = form.save(commit=False)
            instance.slot_number = next_val
            instance.save()
        _log_master_audit(self.request, 'create', 'SlotMaster', str(instance.pk), f'Created SlotMaster: {instance.slot_number}')
        messages.success(self.request, f'Slot {instance.slot_number} created successfully.')
        return redirect(self.success_url)


class SlotMasterUpdateView(UpdateView):
    model = SlotMaster
    fields = ['slot_number', 'slot_name', 'slot_time', 'is_active']
    template_name = 'masters/slot_master_form.html'
    success_url = reverse_lazy('slot_master_manage')
    
    def get_form(self, form_class=None):
        form = super().get_form(form_class)
        if 'slot_number' in form.fields:
            form.fields['slot_number'].disabled = True
        return form
    
    def form_valid(self, form):
        _log_master_audit(self.request, 'update', 'SlotMaster', str(self.object.pk), f'Updated SlotMaster: {self.object.slot_number}')
        messages.success(self.request, 'Slot updated successfully.')
        return super().form_valid(form)


@login_required
@super_admin_required
def slot_master_delete(request, pk):
    """Delete a slot - function-based view to handle GET requests properly"""
    slot = get_object_or_404(SlotMaster, pk=pk)
    slot_number = slot.slot_number
    slot.delete()
    _log_master_audit(request, 'delete', 'SlotMaster', str(pk), f'Deleted SlotMaster: {slot_number}')
    messages.success(request, f'Slot {slot_number} deleted successfully.')
    return redirect('slot_master_manage')


@login_required
@super_admin_required
def slot_master_enable(request, pk):
    """Enable a slot master record."""
    slot = get_object_or_404(SlotMaster, pk=pk)
    slot.is_active = True
    slot.save()
    _log_master_audit(request, 'update', 'SlotMaster', str(pk), f'Enabled SlotMaster: {slot.slot_number}')
    messages.success(request, f'Slot {slot.slot_number} enabled successfully.')
    return redirect('slot_master_manage')


@login_required
@super_admin_required
def slot_master_disable(request, pk):
    """Disable a slot master record."""
    slot = get_object_or_404(SlotMaster, pk=pk)
    slot.is_active = False
    slot.save()
    _log_master_audit(request, 'update', 'SlotMaster', str(pk), f'Disabled SlotMaster: {slot.slot_number}')
    messages.success(request, f'Slot {slot.slot_number} disabled successfully.')
    return redirect('slot_master_manage')

# -----------------------
# Shift Master CRUD Views
# -----------------------
@method_decorator(module_access_required('shift_master_manage'), name='dispatch')
class ShiftMasterManagementView(ListView):
    model = ShiftMaster
    template_name = 'masters/shift_master_manage.html'
    context_object_name = 'shift_list'
    paginate_by = 20

    def get_queryset(self):
        queryset = ShiftMaster.objects.all()

        # Search filter
        search = self.request.GET.get('search', '')
        if search:
            queryset = queryset.filter(shift_name__icontains=search)

        # Status filter
        status = self.request.GET.get('status', '')
        if status == 'active':
            queryset = queryset.filter(is_active=True)
        elif status == 'inactive':
            queryset = queryset.filter(is_active=False)

        return queryset.order_by('shift_name')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['master_type'] = 'shift_master'
        context['master_name'] = 'Shift Master'
        context['search'] = self.request.GET.get('search', '')
        context['status_filter'] = self.request.GET.get('status', '')

        # Template-safety: avoid calling queryset methods with args inside templates.
        # Compute active slot counts once per row for display.
        shift_objs = list(context.get('shift_list', []))
        for s in shift_objs:
            try:
                s.active_slot_count = s.shift_slots.filter(is_active=True).count()
            except Exception:
                s.active_slot_count = 0
        context['shift_list'] = shift_objs
        return context


@method_decorator(module_access_required('shift_master_manage'), name='dispatch')
class ShiftMasterCreateView(CreateView):
    model = ShiftMaster
    fields = ['shift_name', 'is_active']
    template_name = 'masters/shift_master_form.html'
    success_url = reverse_lazy('shift_master_manage')

    def form_valid(self, form):
        response = super().form_valid(form)
        _log_master_audit(self.request, 'create', 'ShiftMaster', str(self.object.pk), f'Created ShiftMaster: {self.object.shift_name}')
        messages.success(self.request, f'Shift "{self.object.shift_name}" created successfully.')
        return response


@method_decorator(module_access_required('shift_master_manage'), name='dispatch')
class ShiftMasterUpdateView(UpdateView):
    model = ShiftMaster
    fields = ['shift_name', 'is_active']
    template_name = 'masters/shift_master_form.html'
    success_url = reverse_lazy('shift_master_manage')

    def form_valid(self, form):
        _log_master_audit(self.request, 'update', 'ShiftMaster', str(self.object.pk), f'Updated ShiftMaster: {self.object.shift_name}')
        messages.success(self.request, f'Shift "{self.object.shift_name}" updated successfully.')
        return super().form_valid(form)


@login_required
@super_admin_required
def shift_master_delete(request, pk):
    """Delete a shift master record."""
    shift = get_object_or_404(ShiftMaster, pk=pk)
    shift_name = shift.shift_name
    shift.delete()
    _log_master_audit(request, 'delete', 'ShiftMaster', str(pk), f'Deleted ShiftMaster: {shift_name}')
    messages.success(request, f'Shift "{shift_name}" deleted successfully.')
    return redirect('shift_master_manage')


@login_required
@super_admin_required
def shift_master_enable(request, pk):
    """Enable a shift master record."""
    shift = get_object_or_404(ShiftMaster, pk=pk)
    shift.is_active = True
    shift.save()
    _log_master_audit(request, 'update', 'ShiftMaster', str(pk), f'Enabled ShiftMaster: {shift.shift_name}')
    messages.success(request, f'Shift "{shift.shift_name}" enabled successfully.')
    return redirect('shift_master_manage')


@login_required
@super_admin_required
def shift_master_disable(request, pk):
    """Disable a shift master record."""
    shift = get_object_or_404(ShiftMaster, pk=pk)
    shift.is_active = False
    shift.save()
    _log_master_audit(request, 'update', 'ShiftMaster', str(pk), f'Disabled ShiftMaster: {shift.shift_name}')
    messages.success(request, f'Shift "{shift.shift_name}" disabled successfully.')
    return redirect('shift_master_manage')


@login_required
@module_access_required('shift_master_manage')
def shift_assign_slots(request, pk):
    """
    Assign SlotMaster records to a ShiftMaster.
    Saves to ShiftSlotMapping, supporting overlap (same slot can be mapped to multiple shifts).
    """
    shift = get_object_or_404(ShiftMaster, pk=pk)

    if request.method == "POST":
        selected_ids_raw = request.POST.getlist("slot_ids")  # list of SlotMaster PKs (strings)
        selected_ids = set()
        for v in selected_ids_raw:
            try:
                selected_ids.add(int(v))
            except (TypeError, ValueError):
                continue

        active_slots_qs = SlotMaster.objects.filter(is_active=True).values_list("pk", flat=True)
        active_slot_ids = set(active_slots_qs)
        # Only allow selections from active slots (fail-safe)
        selected_ids = selected_ids.intersection(active_slot_ids)

        existing_qs = ShiftSlotMapping.objects.filter(shift=shift)
        existing_ids = set(existing_qs.values_list("slot_id", flat=True))

        # Enable selected mappings
        if selected_ids:
            existing_qs.filter(slot_id__in=selected_ids).update(is_active=True)

        # Disable non-selected mappings
        if existing_ids:
            existing_qs.exclude(slot_id__in=selected_ids).update(is_active=False)

        # Create missing mappings
        to_create = list(selected_ids - existing_ids)
        if to_create:
            ShiftSlotMapping.objects.bulk_create(
                [ShiftSlotMapping(shift=shift, slot_id=sid, is_active=True) for sid in to_create]
            )

        _log_master_audit(
            request,
            "update",
            "ShiftMaster",
            str(shift.pk),
            f"Assigned {len(selected_ids)} slot(s) for ShiftMaster {shift.shift_name}",
            additional_data={"selected_slot_ids": sorted(selected_ids)},
        )
        messages.success(request, f"Slots updated for shift \"{shift.shift_name}\".")
        return redirect("shift_assign_slots", pk=shift.pk)

    # GET: render assignment UI
    slot_list = SlotMaster.objects.filter(is_active=True).order_by("slot_number")
    mapped_slot_ids = set(
        ShiftSlotMapping.objects.filter(shift=shift, is_active=True).values_list("slot_id", flat=True)
    )

    return render(
        request,
        "masters/shift_assign_slots.html",
        {
            "shift": shift,
            "slot_list": slot_list,
            "mapped_slot_ids": mapped_slot_ids,
        },
    )

# Ingredient Master CRUD Views
class IngredientMasterManagementView(ListView):
    model = IngredientMaster
    template_name = 'masters/ingredient_master_manage.html'
    context_object_name = 'ingredient_list'
    paginate_by = 20

    def get_queryset(self):
        queryset = IngredientMaster.objects.all()
        
        # Search filter
        search = self.request.GET.get('search', '')
        if search:
            queryset = queryset.filter(
                db_models.Q(ingredient_code__icontains=search) |
                db_models.Q(ingredient_name__icontains=search) |
                db_models.Q(description__icontains=search)
            )
        
        # Status filter
        status = self.request.GET.get('status', '')
        if status == 'active':
            queryset = queryset.filter(is_active=True)
        elif status == 'inactive':
            queryset = queryset.filter(is_active=False)
        
        return queryset.order_by('ingredient_name')
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['master_type'] = 'ingredient'
        context['master_name'] = 'Ingredient Master'
        context['required_fields'] = ['ingredient_name']
        context['optional_fields'] = ['description', 'is_active']
        context['search'] = self.request.GET.get('search', '')
        context['status_filter'] = self.request.GET.get('status', '')
        return context

class IngredientMasterCreateView(CreateView):
    model = IngredientMaster
    fields = ['ingredient_name', 'description', 'is_active']
    template_name = 'masters/ingredient_master_form.html'
    success_url = reverse_lazy('ingredient_master_manage')
    
    def form_valid(self, form):
        # Auto-generate unique ingredient code by incrementing from highest existing code
        # Get all existing ingredient codes that start with ING
        existing_codes = IngredientMaster.objects.filter(
            ingredient_code__startswith='ING'
        ).values_list('ingredient_code', flat=True)
        
        # Extract numbers from existing codes and find the highest
        max_number = 0
        for code in existing_codes:
            try:
                # Extract number part after 'ING'
                number_part = int(code.replace('ING', ''))
                if number_part > max_number:
                    max_number = number_part
            except (ValueError, AttributeError):
                continue
        
        # Increment by 1
        next_number = max_number + 1
        
        # Ensure it doesn't exceed 999
        if next_number > 999:
            raise ValueError("Maximum ingredient code limit reached (ING999)")
        
        # Generate new code
        ingredient_code = f"ING{next_number:03d}"
        
        # Set the auto-generated code
        instance = form.save(commit=False)
        instance.ingredient_code = ingredient_code
        instance.save()
        _log_master_audit(self.request, 'create', 'IngredientMaster', str(instance.pk), f'Created IngredientMaster: {instance.ingredient_name} ({ingredient_code})')
        messages.success(self.request, f'Ingredient "{instance.ingredient_name}" ({ingredient_code}) created successfully.')
        return redirect(self.success_url)

class IngredientMasterUpdateView(UpdateView):
    model = IngredientMaster
    fields = ['ingredient_name', 'description', 'is_active']
    template_name = 'masters/ingredient_master_form.html'
    success_url = reverse_lazy('ingredient_master_manage')
    
    def form_valid(self, form):
        # Preserve the existing ingredient_code (it's auto-generated and shouldn't be changed)
        instance = form.save(commit=False)
        if not instance.ingredient_code:
            # If somehow code is missing, generate one (shouldn't happen)
            existing_codes = IngredientMaster.objects.filter(
                ingredient_code__startswith='ING'
            ).values_list('ingredient_code', flat=True)
            max_number = 0
            for code in existing_codes:
                try:
                    number_part = int(code.replace('ING', ''))
                    if number_part > max_number:
                        max_number = number_part
                except (ValueError, AttributeError):
                    continue
            instance.ingredient_code = f"ING{max_number + 1:03d}"
        instance.save()
        _log_master_audit(self.request, 'update', 'IngredientMaster', str(self.object.pk), f'Updated IngredientMaster: {instance.ingredient_name}')
        messages.success(self.request, f'Ingredient "{instance.ingredient_name}" updated successfully.')
        return redirect(self.success_url)

@login_required
@super_admin_required
def ingredient_master_delete(request, pk):
    """Delete an ingredient - function-based view to handle GET requests properly"""
    ingredient = get_object_or_404(IngredientMaster, pk=pk)
    ingredient_name = ingredient.ingredient_name
    ingredient.delete()
    _log_master_audit(request, 'delete', 'IngredientMaster', str(pk), f'Deleted IngredientMaster: {ingredient_name}')
    messages.success(request, f'Ingredient {ingredient_name} deleted successfully.')
    return redirect('ingredient_master_manage')

@login_required
@super_admin_required
def ingredient_master_enable(request, pk):
    """Enable an ingredient"""
    ingredient = get_object_or_404(IngredientMaster, pk=pk)
    ingredient.is_active = True
    ingredient.save()
    _log_master_audit(request, 'update', 'IngredientMaster', str(pk), f'Enabled IngredientMaster: {ingredient.ingredient_name}')
    messages.success(request, f'Ingredient {ingredient.ingredient_name} enabled successfully.')
    return redirect('ingredient_master_manage')

@login_required
@super_admin_required
def ingredient_master_disable(request, pk):
    """Disable an ingredient"""
    ingredient = get_object_or_404(IngredientMaster, pk=pk)
    ingredient.is_active = False
    ingredient.save()
    _log_master_audit(request, 'update', 'IngredientMaster', str(pk), f'Disabled IngredientMaster: {ingredient.ingredient_name}')
    messages.success(request, f'Ingredient {ingredient.ingredient_name} disabled successfully.')
    return redirect('ingredient_master_manage')

# Treatment Room Master CRUD Views
@method_decorator(login_required, name='dispatch')
@method_decorator(super_admin_required, name='dispatch')
class TreatmentRoomManagementView(ListView):
    model = TreatmentRoom
    template_name = 'masters/treatment_room_master_manage.html'
    context_object_name = 'treatment_room_list'
    paginate_by = 20

    def get_queryset(self):
        queryset = TreatmentRoom.objects.select_related('room_name', 'floor_master').all()

        search = self.request.GET.get('search', '')
        if search:
            queryset = queryset.filter(
                db_models.Q(room_no__icontains=search) |
                db_models.Q(tr_code__icontains=search) |
                db_models.Q(room_name__name__icontains=search) |
                db_models.Q(floor_master__name__icontains=search) |
                db_models.Q(description__icontains=search)
            )

        status = self.request.GET.get('status', '')
        if status == 'active':
            queryset = queryset.filter(is_active=True)
        elif status == 'inactive':
            queryset = queryset.filter(is_active=False)

        return queryset.order_by('room_no')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['search'] = self.request.GET.get('search', '')
        context['status_filter'] = self.request.GET.get('status', '')
        context['master_type'] = 'treatmentroom'
        context['master_name'] = 'Treatment Room Master'
        return context


class TreatmentRoomForm(forms.ModelForm):
    class Meta:
        model = TreatmentRoom
        fields = ['room_no', 'room_name', 'floor_master', 'description', 'is_active']
        widgets = {
            'room_no': forms.TextInput(attrs={'class': 'form-control', 'placeholder': 'e.g., TR-1, RM-101'}),
            'room_name': forms.Select(attrs={'class': 'form-select'}),
            'floor_master': forms.Select(attrs={'class': 'form-select'}),
            'description': forms.Textarea(attrs={'class': 'form-control', 'rows': 3}),
        }
        labels = {
            'room_no': 'Room Number',
            'room_name': 'Room Name',
            'floor_master': 'Floor',
        }

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields['room_name'].queryset = RoomName.objects.filter(is_active=True)
        self.fields['floor_master'].queryset = Floor.objects.filter(is_active=True)
        self.fields['room_no'].required = True
        self.fields['room_name'].required = False
        self.fields['floor_master'].required = False


@method_decorator(login_required, name='dispatch')
@method_decorator(super_admin_required, name='dispatch')
class TreatmentRoomCreateView(CreateView):
    model = TreatmentRoom
    form_class = TreatmentRoomForm
    template_name = 'masters/treatment_room_master_form.html'
    success_url = reverse_lazy('treatment_room_master_manage')

    def form_valid(self, form):
        instance = form.save(commit=False)
        instance.save()
        _log_master_audit(self.request, 'create', 'TreatmentRoom', str(instance.pk), f'Created TreatmentRoom: {instance.room_no}')
        messages.success(self.request, f'Treatment room {instance.room_no} created successfully.')
        return redirect(self.success_url)


@method_decorator(login_required, name='dispatch')
@method_decorator(super_admin_required, name='dispatch')
class TreatmentRoomUpdateView(UpdateView):
    model = TreatmentRoom
    form_class = TreatmentRoomForm
    template_name = 'masters/treatment_room_master_form.html'
    success_url = reverse_lazy('treatment_room_master_manage')

    def get_form(self, form_class=None):
        form = super().get_form(form_class)
        form.fields['room_no'].disabled = True
        return form

    def form_valid(self, form):
        _log_master_audit(self.request, 'update', 'TreatmentRoom', str(self.object.pk), f'Updated TreatmentRoom: {self.object.room_no}')
        messages.success(self.request, 'Treatment room updated successfully.')
        return super().form_valid(form)

@login_required
@super_admin_required
def treatment_room_master_delete(request, pk):
    """Delete a treatment room - function-based view to handle GET requests properly"""
    room = get_object_or_404(TreatmentRoom, pk=pk)
    room_no = room.room_no
    room.delete()
    _log_master_audit(request, 'delete', 'TreatmentRoom', str(pk), f'Deleted TreatmentRoom: {room_no}')
    messages.success(request, f'Treatment room {room_no} deleted successfully.')
    return redirect('treatment_room_master_manage')

@login_required
@super_admin_required
def treatment_room_master_enable(request, pk):
    """Enable a treatment room"""
    room = get_object_or_404(TreatmentRoom, pk=pk)
    room.is_active = True
    room.save()
    _log_master_audit(request, 'update', 'TreatmentRoom', str(pk), f'Enabled TreatmentRoom: {room.room_no}')
    messages.success(request, f'Treatment room {room.room_no} enabled successfully.')
    return redirect('treatment_room_master_manage')

@login_required
@super_admin_required
def treatment_room_master_disable(request, pk):
    """Disable a treatment room"""
    room = get_object_or_404(TreatmentRoom, pk=pk)
    room.is_active = False
    room.save()
    _log_master_audit(request, 'update', 'TreatmentRoom', str(pk), f'Disabled TreatmentRoom: {room.room_no}')
    messages.success(request, f'Treatment room {room.room_no} disabled successfully.')
    return redirect('treatment_room_master_manage')


# Therapist Master CRUD Views
class TherapistMasterManagementView(ListView):
    model = TherapistMaster
    template_name = 'masters/therapist_master_manage.html'
    context_object_name = 'therapist_list'
    paginate_by = 20

    def get_queryset(self):
        # Get status filter FIRST
        status = self.request.GET.get('status', '').strip().lower()
        
        # Start with base queryset and apply status filter immediately
        if status == 'active':
            queryset = TherapistMaster.objects.filter(is_active=True)
        elif status == 'inactive':
            queryset = TherapistMaster.objects.filter(is_active=False)
        else:
            # No status filter - show all
            queryset = TherapistMaster.objects.all()
        
        # Apply prefetch_related after filtering
        queryset = queryset.prefetch_related('leaves', 'therapist_grade')
        
        # Search filter
        search = self.request.GET.get('search', '').strip()
        if search:
            queryset = queryset.filter(
                db_models.Q(therapist_code__icontains=search) |
                db_models.Q(therapist_name__icontains=search) |
                db_models.Q(employee_code__icontains=search)
            )
        
        # Grade filter
        grade = self.request.GET.get('grade', '')
        if grade:
            queryset = queryset.filter(therapist_grade__id=grade)
        
        return queryset.order_by('therapist_name')
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['grades'] = TherapistGrade.objects.filter(is_active=True).order_by('display_order', 'name')
        context['search'] = self.request.GET.get('search', '')
        context['status_filter'] = self.request.GET.get('status', '')
        context['grade_filter'] = self.request.GET.get('grade', '')
        # Bulk upload context
        # Note: Therapist master has complex relationships, basic bulk upload support
        context['master_type'] = 'therapist'
        context['master_name'] = 'Therapist Master'
        context['required_fields'] = ['therapist_code', 'therapist_name']
        context['optional_fields'] = ['employee_code', 'therapist_gender', 'is_bathing_therapist', 'is_active']
        return context

class TherapistMasterCreateView(CreateView):
    model = TherapistMaster
    fields = ['therapist_name', 'therapist_grade', 'employee_code', 'therapist_gender', 'is_bathing_therapist', 'is_active']
    template_name = 'masters/therapist_master_form.html'
    success_url = reverse_lazy('therapist_master_manage')
    
    def get_form(self, form_class=None):
        form = super().get_form(form_class)
        # Make therapist_code read-only/hidden since it's auto-generated
        if 'therapist_code' in form.fields:
            form.fields['therapist_code'].widget.attrs['readonly'] = True
        # Update grade field to use TherapistGrade queryset
        if 'therapist_grade' in form.fields:
            form.fields['therapist_grade'].queryset = TherapistGrade.objects.filter(is_active=True).order_by('display_order', 'name')
        # Customize gender field to use Select widget
        if 'therapist_gender' in form.fields:
            form.fields['therapist_gender'].widget = forms.Select(choices=[('', 'Select Gender'), ('Male', 'Male'), ('Female', 'Female')])
            form.fields['therapist_gender'].widget.attrs['class'] = 'form-select'
            form.fields['therapist_gender'].initial = 'Male'
        return form
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['grades'] = TherapistGrade.objects.filter(is_active=True).order_by('display_order', 'name')
        context['available_shifts'] = ShiftMaster.objects.filter(is_active=True).order_by('shift_name')
        context['selected_shift_ids'] = []
        return context
    
    def form_valid(self, form):
        """Create user account and profile when creating therapist master"""
        from authentication.models import TherapistProfile
        
        # Auto-generate unique therapist code by checking both tables
        # Get existing codes from TherapistMaster
        existing_master_codes = list(TherapistMaster.objects.filter(
            therapist_code__startswith='THP'
        ).values_list('therapist_code', flat=True))
        
        # Get existing codes from TherapistProfile
        existing_profile_codes = list(TherapistProfile.objects.filter(
            therapist_code__startswith='THP'
        ).values_list('therapist_code', flat=True))
        
        existing_codes = set(existing_master_codes + existing_profile_codes)
        
        # Extract numbers from existing codes and find the highest
        max_number = 0
        for code in existing_codes:
            try:
                # Extract number part after 'THP'
                number_part = int(code.replace('THP', ''))
                if number_part > max_number:
                    max_number = number_part
            except (ValueError, AttributeError):
                continue
        
        # Increment by 1
        next_number = max_number + 1
        
        # Ensure it doesn't exceed 9999
        if next_number > 9999:
            raise ValueError("Maximum therapist code limit reached (THP9999)")
        
        # Generate new code
        therapist_code = f"THP{next_number:04d}"
        
        # Set the auto-generated code
        instance = form.save(commit=False)
        instance.therapist_code = therapist_code
        
        therapist_name = form.cleaned_data['therapist_name']
        therapist_grade = form.cleaned_data.get('therapist_grade')
        employee_code = form.cleaned_data['employee_code']
        is_bathing_therapist = form.cleaned_data.get('is_bathing_therapist', True)
        
        # Generate username and email
        username = therapist_code.lower().replace(' ', '_')
        email = f"{username}@avcri.local"
        
        # Check if username exists and guarantee uniqueness
        from authentication.models import CustomUser
        base_username = username
        counter = 1
        while CustomUser.objects.filter(username=username).exists():
            username = f"{base_username}_{counter}"
            counter += 1
        
        # Split name
        name_parts = therapist_name.split(maxsplit=1)
        first_name = name_parts[0] if name_parts else therapist_name
        last_name = name_parts[1] if len(name_parts) > 1 else ''
        
        # Create user account
        default_password = f"{therapist_code}@{employee_code}"
        user = CustomUser.objects.create_user(
            username=username,
            email=email,
            password=default_password,
            first_name=first_name,
            last_name=last_name,
            role='therapist'
        )
        
        # Create therapist profile
        from authentication.models import TherapistProfile
        # Get grade name for profile (if grade is a ForeignKey)
        grade_name = therapist_grade.name if therapist_grade else 'Junior'
        therapist_profile = TherapistProfile.objects.create(
            user=user,
            therapist_code=therapist_code,
            employee_code=employee_code,
            therapist_grade=grade_name,
            is_bathing_therapist=is_bathing_therapist
        )
        
        # Link user and profile to the therapist master instance
        instance.user = user
        instance.therapist_profile = therapist_profile
        instance.save()

        # Handle therapist -> shift mapping
        selected_shift_ids = set()
        for sid in self.request.POST.getlist('shift_ids'):
            try:
                selected_shift_ids.add(int(sid))
            except (TypeError, ValueError):
                continue
        if selected_shift_ids:
            valid_shift_ids = set(
                ShiftMaster.objects.filter(id__in=selected_shift_ids).values_list('id', flat=True)
            )
            TherapistShiftMapping.objects.bulk_create(
                [TherapistShiftMapping(therapist=instance, shift_id=sid, is_active=True) for sid in valid_shift_ids]
            )

        _log_master_audit(self.request, 'create', 'TherapistMaster', str(instance.pk), f'Created TherapistMaster: {therapist_name} ({therapist_code})')
        messages.success(self.request, f'Therapist "{therapist_name}" ({therapist_code}) created successfully. Username: {username}, Default password: {default_password}')
        return redirect(self.success_url)

class TherapistMasterUpdateView(UpdateView):
    model = TherapistMaster
    fields = ['therapist_name', 'therapist_grade', 'employee_code', 'therapist_gender', 'is_bathing_therapist', 'is_active']
    template_name = 'masters/therapist_master_form.html'
    success_url = reverse_lazy('therapist_master_manage')
    
    def get_form(self, form_class=None):
        form = super().get_form(form_class)
        # Make therapist_code read-only since it shouldn't be changed
        if 'therapist_code' in form.fields:
            form.fields['therapist_code'].widget.attrs['readonly'] = True
        # Update grade field to use TherapistGrade queryset
        if 'therapist_grade' in form.fields:
            form.fields['therapist_grade'].queryset = TherapistGrade.objects.filter(is_active=True).order_by('display_order', 'name')
        # Customize gender field to use Select widget
        if 'therapist_gender' in form.fields:
            form.fields['therapist_gender'].widget = forms.Select(choices=[('', 'Select Gender'), ('Male', 'Male'), ('Female', 'Female')])
            form.fields['therapist_gender'].widget.attrs['class'] = 'form-select'
        return form
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['grades'] = TherapistGrade.objects.filter(is_active=True).order_by('display_order', 'name')
        context['available_shifts'] = ShiftMaster.objects.filter(is_active=True).order_by('shift_name')
        context['selected_shift_ids'] = list(
            TherapistShiftMapping.objects.filter(therapist=self.object, is_active=True).values_list('shift_id', flat=True)
        )
        
        # Add leave management context
        context['therapist_leaves'] = [
            (0, 'Monday'), (1, 'Tuesday'), (2, 'Wednesday'), (3, 'Thursday'),
            (4, 'Friday'), (5, 'Saturday'), (6, 'Sunday')
        ]
        
        # Get existing leave days for this therapist
        existing_leaves = TherapistLeaveMaster.objects.filter(
            therapist=self.object, is_active=True
        ).values_list('leave_day', flat=True)
        context['existing_leaves'] = list(existing_leaves)
        
        return context
    
    def form_valid(self, form):
        # Save the therapist first
        response = super().form_valid(form)
        _log_master_audit(self.request, 'update', 'TherapistMaster', str(self.object.pk), f'Updated TherapistMaster: {self.object.therapist_name} ({self.object.therapist_code})')

        # Handle therapist -> shift mapping
        selected_shift_ids = set()
        for sid in self.request.POST.getlist('shift_ids'):
            try:
                selected_shift_ids.add(int(sid))
            except (TypeError, ValueError):
                continue

        existing_qs = TherapistShiftMapping.objects.filter(therapist=self.object)
        existing_ids = set(existing_qs.values_list('shift_id', flat=True))
        valid_shift_ids = set(
            ShiftMaster.objects.filter(id__in=selected_shift_ids).values_list('id', flat=True)
        )

        if valid_shift_ids:
            existing_qs.filter(shift_id__in=valid_shift_ids).update(is_active=True)
        if existing_ids:
            existing_qs.exclude(shift_id__in=valid_shift_ids).update(is_active=False)

        to_create = list(valid_shift_ids - existing_ids)
        if to_create:
            TherapistShiftMapping.objects.bulk_create(
                [TherapistShiftMapping(therapist=self.object, shift_id=sid, is_active=True) for sid in to_create]
            )
        
        # Handle leave days
        if 'leave_days' in self.request.POST:
            selected_days = self.request.POST.getlist('leave_days')
            
            # Remove existing leave records for this therapist
            TherapistLeaveMaster.objects.filter(therapist=self.object).delete()
            
            # Create new leave records for selected days
            for day in selected_days:
                TherapistLeaveMaster.objects.create(
                    therapist=self.object,
                    leave_day=int(day),
                    is_active=True
                )
        else:
            # If no days selected, remove all leave records
            TherapistLeaveMaster.objects.filter(therapist=self.object).delete()
        
        messages.success(self.request, 'Therapist and leave schedule updated successfully.')
        return response

@login_required
@super_admin_required
def therapist_master_export_excel(request):
    """Export all therapist master data to Excel"""
    try:
        # Check if openpyxl is available
        try:
            import openpyxl
        except ImportError:
            messages.error(request, 'openpyxl is required for Excel export. Please install it: pip install openpyxl')
            return redirect('therapist_master_manage')
        
        # Get all therapists with related data
        therapists = TherapistMaster.objects.select_related(
            'therapist_grade', 'user', 'therapist_profile'
        ).prefetch_related('leaves').all().order_by('therapist_name')
        
        # Prepare data for Excel
        data = []
        for therapist in therapists:
            # Get leave days as comma-separated string
            leave_days = ', '.join([leave.leave_day_name for leave in therapist.leaves.all()]) if therapist.leaves.exists() else 'None'
            
            # Get user info if available
            username = therapist.user.username if therapist.user else 'N/A'
            email = therapist.user.email if therapist.user else 'N/A'
            
            data.append({
                'Therapist Code': therapist.therapist_code,
                'Therapist Name': therapist.therapist_name,
                'Grade': therapist.therapist_grade.name if therapist.therapist_grade else 'N/A',
                'Employee Code': therapist.employee_code,
                'Gender': therapist.therapist_gender or 'N/A',
                'Bathing Therapist': 'Yes' if therapist.is_bathing_therapist else 'No',
                'Leave Days': leave_days,
                'Status': 'Active' if therapist.is_active else 'Inactive',
                'Username': username,
                'Email': email,
                'Created At': therapist.created_at.strftime('%Y-%m-%d %H:%M:%S') if therapist.created_at else '',
                'Updated At': therapist.updated_at.strftime('%Y-%m-%d %H:%M:%S') if therapist.updated_at else '',
            })
        
        # Create DataFrame
        df = pd.DataFrame(data)
        
        # Create Excel file in memory
        output = BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Therapist Master', index=False)
            
            # Get the workbook and worksheet
            workbook = writer.book
            worksheet = writer.sheets['Therapist Master']
            
            # Auto-adjust column widths
            from openpyxl.utils import get_column_letter
            for idx, col in enumerate(df.columns, 1):
                max_length = max(
                    df[col].astype(str).map(len).max(),
                    len(col)
                ) + 2
                col_letter = get_column_letter(idx)
                worksheet.column_dimensions[col_letter].width = min(max_length, 50)
        
        # Prepare response
        output.seek(0)
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f'therapist_master_export_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        return response
        
    except Exception as e:
        messages.error(request, f'Error exporting data: {str(e)}')
        import traceback
        traceback.print_exc()
        return redirect('therapist_master_manage')

@login_required
@super_admin_required
def download_master_template(request, master_type):
    """Download example Excel template for bulk upload"""
    try:
        master_type = master_type.lower()
        
        # Define template data for each master type
        template_data = {}
        
        if master_type == 'therapy' or master_type == 'therapymaster':
            template_data = {
                'columns': ['therapy_code', 'therapy_name', 'therapy_group', 'therapy_category', 
                           'therapy_duration', 'bathing_required', 'treatment_room_required', 
                           'cost_of_therapy', 'patient_type', 'no_of_super_senior_therapist',
                           'no_of_senior_therapist', 'no_of_junior_therapist', 'no_of_trainee_therapist', 'is_active'],
                'example_rows': [
                    ['THEB001', 'Abhyanga', 'Panchakarma', 'Medium', '45 min', 'Yes', 'Yes', '500.00', 'IP', 1, 0, 0, 0, 'true'],
                    ['THEB002', 'Dhara', 'Panchakarma', 'Hard', '60 min', 'Yes', 'Yes', '800.00', 'IP', 0, 1, 1, 0, 'true'],
                ]
            }
        elif master_type == 'item' or master_type == 'itemmaster':
            template_data = {
                'columns': ['item_name', 'item_group', 'umo', 'min_quantity', 'max_quantity', 'procured_from', 'is_active'],
                'example_rows': [
                    ['Sesame Oil', 'Oil', 'ml', 10, 100, 'IP Pharmacy', 'true'],
                    ['Turmeric Powder', 'Herb', 'gm', 5, 50, 'Kashaya Room', 'true'],
                ]
            }
        elif master_type == 'slot' or master_type == 'slotmaster':
            template_data = {
                'columns': ['slot_number', 'slot_time', 'slot_name', 'is_active'],
                'example_rows': [
                    [1, '08:00:00', 'Morning Slot', 'true'],
                    [2, '10:00:00', 'Mid Morning Slot', 'true'],
                ]
            }
        elif master_type == 'ingredient' or master_type == 'ingredientmaster':
            template_data = {
                'columns': ['ingredient_name', 'description', 'is_active'],
                'example_rows': [
                    ['Turmeric', 'Anti-inflammatory herb', 'true'],
                    ['Ginger', 'Digestive aid', 'true'],
                ]
            }
        elif master_type == 'treatmentroom' or master_type == 'treatment_room':
            template_data = {
                'columns': ['room_no', 'floor', 'description', 'tr_code', 'is_active'],
                'example_rows': [
                    ['TR001', 'Ground Floor', 'Therapy Room 1', 'TR001', 'true'],
                    ['TR002', 'First Floor', 'Therapy Room 2', 'TR002', 'true'],
                ]
            }
        elif master_type == 'procured':
            template_data = {
                'columns': ['name', 'is_active'],
                'example_rows': [
                    ['IP Pharmacy', 'true'],
                    ['Kashaya Room', 'true'],
                ]
            }
        elif master_type == 'oilclass' or master_type == 'oil_class':
            template_data = {
                'columns': ['name', 'is_active'],
                'example_rows': [
                    ['Sesame Oil (Tila Taila)', 'true'],
                    ['Coconut Oil (Narikela Taila)', 'true'],
                ]
            }
        elif master_type == 'usagetype' or master_type == 'usage_type':
            template_data = {
                'columns': ['name', 'is_active'],
                'example_rows': [
                    ['Internal (IP Use)', 'true'],
                    ['Dispensed (OP Use)', 'true'],
                ]
            }
        elif master_type == 'prescriptiontime' or master_type == 'prescription_time':
            template_data = {
                'columns': ['name', 'is_active'],
                'example_rows': [
                    ['8 AM', 'true'],
                    ['12 PM', 'true'],
                ]
            }
        elif master_type == 'prescriptionscope' or master_type == 'prescription_scope':
            template_data = {
                'columns': ['name', 'is_active'],
                'example_rows': [
                    ['Internal', 'true'],
                    ['External', 'true'],
                ]
            }
        elif master_type == 'unitno' or master_type == 'unit_no':
            template_data = {
                'columns': ['unit_name', 'is_active'],
                'example_rows': [
                    ['ml', 'true'],
                    ['gm', 'true'],
                ]
            }
        elif master_type == 'nursetreatment' or master_type == 'nurse_treatment':
            template_data = {
                'columns': ['treatment_code', 'name', 'number_of_nurse_required', 'is_active'],
                'example_rows': [
                    ['NT001', 'Nurse Treatment 1', 2, 'true'],
                    ['NT002', 'Nurse Treatment 2', 1, 'true'],
                ]
            }
        elif master_type == 'nursemaster' or master_type == 'nurse_master':
            template_data = {
                'columns': ['nurse_code', 'nurse_name', 'designation', 'is_active'],
                'example_rows': [
                    ['NUR001', 'Nurse Name 1', 'Senior Nurse', 'true'],
                    ['NUR002', 'Nurse Name 2', 'Staff Nurse', 'true'],
                ]
            }
        elif master_type == 'prescriptionusage' or master_type == 'prescription_usage':
            template_data = {
                'columns': ['name', 'is_active'],
                'example_rows': [
                    ['Internal Use', 'true'],
                    ['External Use', 'true'],
                ]
            }
        elif master_type == 'therapist' or master_type == 'therapistmaster':
            template_data = {
                'columns': ['therapist_code', 'therapist_name', 'employee_code', 'therapist_gender', 'is_bathing_therapist', 'is_active'],
                'example_rows': [
                    ['THR001', 'Therapist Name 1', 'EMP001', 'Male', 'true', 'true'],
                    ['THR002', 'Therapist Name 2', 'EMP002', 'Female', 'false', 'true'],
                ]
            }
        else:
            messages.error(request, f'Unknown master type: {master_type}')
            return redirect('master_management_dashboard')
        
        # Create DataFrame
        df = pd.DataFrame(template_data['example_rows'], columns=template_data['columns'])
        
        # Create Excel file in memory
        output = BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Template', index=False)
            
            # Get the workbook and worksheet
            workbook = writer.book
            worksheet = writer.sheets['Template']
            
            # Auto-adjust column widths
            from openpyxl.utils import get_column_letter
            from openpyxl.styles import Font, PatternFill, Alignment
            
            # Style header row
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")
            
            for idx, col in enumerate(template_data['columns'], 1):
                col_letter = get_column_letter(idx)
                cell = worksheet[f'{col_letter}1']
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
                
                # Set column width
                max_length = max(
                    df[col].astype(str).map(len).max() if len(df) > 0 else 0,
                    len(col)
                ) + 2
                worksheet.column_dimensions[col_letter].width = min(max_length, 50)
        
        # Prepare response
        output.seek(0)
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f'{master_type}_template.xlsx'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        return response
        
    except Exception as e:
        messages.error(request, f'Error generating template: {str(e)}')
        import traceback
        traceback.print_exc()
        return redirect('master_management_dashboard')

@login_required
@super_admin_required
@require_http_methods(["POST"])
def bulk_upload_master(request):
    """Bulk upload master data from Excel file"""
    try:
        if 'excel_file' not in request.FILES:
            return JsonResponse({'success': False, 'error': 'No file uploaded'}, status=400)
        
        excel_file = request.FILES['excel_file']
        master_type = request.POST.get('master_type', '').lower()
        
        # Validate file extension
        if not excel_file.name.endswith(('.xlsx', '.xls')):
            return JsonResponse({'success': False, 'error': 'Invalid file format. Please upload .xlsx or .xls file.'}, status=400)
        
        # Read Excel file
        try:
            df = pd.read_excel(excel_file, sheet_name=0)
        except Exception as e:
            return JsonResponse({'success': False, 'error': f'Error reading Excel file: {str(e)}'}, status=400)
        
        if df.empty:
            return JsonResponse({'success': False, 'error': 'Excel file is empty'}, status=400)
        
        # Normalize column names (lowercase, strip spaces, replace spaces/underscores)
        df.columns = df.columns.str.lower().str.strip().str.replace(' ', '_').str.replace('-', '_')
        
        created_count = 0
        updated_count = 0
        error_count = 0
        errors = []
        
        # Process based on master type
        with transaction.atomic():
            if master_type == 'therapymaster' or master_type == 'therapy':
                # Required columns: all except therapy_group
                required_cols = ['therapy_code', 'therapy_name', 'therapy_category', 'therapy_duration',
                                 'bathing_required', 'treatment_room_required', 'cost_of_therapy', 'patient_type']
                missing_cols = [c for c in required_cols if c not in df.columns]
                if missing_cols:
                    return JsonResponse({'success': False, 'error': f'Required columns missing: {", ".join(missing_cols)}'}, status=400)
                
                for idx, row in df.iterrows():
                    try:
                        therapy_code = str(row.get('therapy_code', '')).strip()
                        therapy_name = str(row.get('therapy_name', '')).strip()
                        therapy_category = str(row.get('therapy_category', '')).strip()
                        therapy_duration = str(row.get('therapy_duration', '')).strip()
                        bathing_required = str(row.get('bathing_required', '')).strip()
                        treatment_room_required = str(row.get('treatment_room_required', '')).strip()
                        cost_val = row.get('cost_of_therapy')
                        cost_provided = cost_val is not None and cost_val != ''
                        patient_type = str(row.get('patient_type', '')).strip()
                        
                        if not therapy_code or not therapy_name or not therapy_category or not therapy_duration or not bathing_required or not treatment_room_required or not cost_provided or not patient_type:
                            missing = []
                            if not therapy_code: missing.append('therapy_code')
                            if not therapy_name: missing.append('therapy_name')
                            if not therapy_category: missing.append('therapy_category')
                            if not therapy_duration: missing.append('therapy_duration')
                            if not bathing_required: missing.append('bathing_required')
                            if not treatment_room_required: missing.append('treatment_room_required')
                            if not cost_provided: missing.append('cost_of_therapy')
                            if not patient_type: missing.append('patient_type')
                            error_count += 1
                            errors.append(f'Row {idx+2}: Missing required fields ({", ".join(missing)})')
                            continue
                        
                        # Parse boolean fields
                        is_active = _parse_boolean(row.get('is_active', True))
                        
                        # Handle therapy_group as ForeignKey
                        therapy_group_name = str(row.get('therapy_group', '')).strip()
                        therapy_group_obj = None
                        if therapy_group_name:
                            try:
                                therapy_group_obj = TherapyGroup.objects.get(name=therapy_group_name)
                            except TherapyGroup.DoesNotExist:
                                pass
                        
                        # Get or create
                        therapy, created = TherapyMaster.objects.get_or_create(
                            therapy_code=therapy_code,
                            defaults={
                                'therapy_name': therapy_name,
                                'therapy_group': therapy_group_obj,
                                'therapy_category': therapy_category,
                                'therapy_duration': therapy_duration,
                                'bathing_required': bathing_required,
                                'treatment_room_required': treatment_room_required,
                                'cost_of_therapy': _parse_decimal(row.get('cost_of_therapy')) or 0,
                                'patient_type': patient_type if patient_type in ('IP', 'OP') else 'IP',
                                'no_of_super_senior_therapist': _parse_int(row.get('no_of_super_senior_therapist', 0)),
                                'no_of_senior_therapist': _parse_int(row.get('no_of_senior_therapist', 0)),
                                'no_of_junior_therapist': _parse_int(row.get('no_of_junior_therapist', 0)),
                                'no_of_trainee_therapist': _parse_int(row.get('no_of_trainee_therapist', 0)),
                                'total_no_of_therapist': _parse_int(row.get('total_no_of_therapist', 0)),
                                'is_active': is_active,
                            }
                        )
                        
                        if not created:
                            # Update existing
                            therapy.therapy_name = therapy_name
                            therapy.therapy_group = therapy_group_obj
                            therapy.therapy_category = therapy_category
                            therapy.therapy_duration = therapy_duration
                            therapy.bathing_required = bathing_required
                            therapy.treatment_room_required = treatment_room_required
                            therapy.cost_of_therapy = _parse_decimal(row.get('cost_of_therapy')) or 0
                            therapy.patient_type = patient_type if patient_type in ('IP', 'OP') else 'IP'
                            therapy.no_of_super_senior_therapist = _parse_int(row.get('no_of_super_senior_therapist', 0))
                            therapy.no_of_senior_therapist = _parse_int(row.get('no_of_senior_therapist', 0))
                            therapy.no_of_junior_therapist = _parse_int(row.get('no_of_junior_therapist', 0))
                            therapy.no_of_trainee_therapist = _parse_int(row.get('no_of_trainee_therapist', 0))
                            therapy.total_no_of_therapist = _parse_int(row.get('total_no_of_therapist', 0))
                            therapy.is_active = is_active
                            therapy.save()
                            updated_count += 1
                        else:
                            created_count += 1
                    except Exception as e:
                        error_count += 1
                        errors.append(f'Row {idx+2}: {str(e)}')
            
            elif master_type == 'itemmaster' or master_type == 'item':
                if 'item_name' not in df.columns:
                    return JsonResponse({'success': False, 'error': 'Required column missing: item_name'}, status=400)
                
                for idx, row in df.iterrows():
                    try:
                        item_name = str(row.get('item_name', '')).strip()
                        
                        if not item_name:
                            error_count += 1
                            errors.append(f'Row {idx+2}: Missing required field (item_name)')
                            continue
                        
                        is_active = _parse_boolean(row.get('is_active', True))
                        
                        # Handle item_group as ForeignKey
                        item_group_name = str(row.get('item_group', '')).strip()
                        item_group_obj = None
                        if item_group_name:
                            try:
                                item_group_obj = ItemGroup.objects.get(name=item_group_name)
                            except ItemGroup.DoesNotExist:
                                pass
                        
                        # Handle umo as ForeignKey
                        umo_name = str(row.get('umo', '')).strip()
                        umo_obj = None
                        if umo_name:
                            try:
                                umo_obj = UnitNo.objects.get(unit_name=umo_name)
                            except UnitNo.DoesNotExist:
                                pass
                        
                        # Handle procured_from as ForeignKey
                        procured_from_name = str(row.get('procured_from', '')).strip()
                        procured_from_obj = None
                        if procured_from_name:
                            try:
                                procured_from_obj = Procured.objects.get(name=procured_from_name)
                            except Procured.DoesNotExist:
                                pass
                        
                        # Logic: Check if item_code provided exists
                        input_code = str(row.get('item_code', '')).strip()
                        item = None
                        if input_code:
                            try:
                                item = ItemMaster.objects.get(item_code=input_code)
                            except ItemMaster.DoesNotExist:
                                item = None
                                
                        if item:
                            # Update existing
                            item.item_name = item_name
                            item.item_group = item_group_obj
                            item.umo = umo_obj
                            item.min_quantity = _parse_int(row.get('min_quantity'), None)
                            item.max_quantity = _parse_int(row.get('max_quantity'), None)
                            item.procured_from = procured_from_obj
                            item.is_active = is_active
                            item.save()
                            updated_count += 1
                        else:
                            # Create new with AUTO GEN code (NUMERIC)
                            new_code = generate_numeric_code(ItemMaster, 'item_code')
                            ItemMaster.objects.create(
                                item_code=new_code,
                                item_name=item_name,
                                item_group=item_group_obj,
                                umo=umo_obj,
                                min_quantity=_parse_int(row.get('min_quantity'), None),
                                max_quantity=_parse_int(row.get('max_quantity'), None),
                                procured_from=procured_from_obj,
                                is_active=is_active
                            )
                            created_count += 1

                    except Exception as e:
                        error_count += 1
                        errors.append(f'Row {idx+2}: {str(e)}')

            
            elif master_type == 'slotmaster' or master_type == 'slot':
                if 'slot_time' not in df.columns:
                    return JsonResponse({'success': False, 'error': 'Required column missing: slot_time'}, status=400)
                
                for idx, row in df.iterrows():
                    try:
                        input_slot_number = _parse_int(row.get('slot_number'))
                        slot_time_str = str(row.get('slot_time', '')).strip()
                        
                        if not slot_time_str:
                            error_count += 1
                            errors.append(f'Row {idx+2}: Missing required field (slot_time)')
                            continue
                        
                        # Parse time
                        try:
                            from datetime import datetime as dt
                            if isinstance(slot_time_str, str):
                                slot_time = dt.strptime(slot_time_str, '%H:%M:%S').time()
                            else:
                                slot_time = slot_time_str
                        except:
                            try:
                                slot_time = dt.strptime(slot_time_str, '%H:%M').time()
                            except:
                                error_count += 1
                                errors.append(f'Row {idx+2}: Invalid time format for slot_time')
                                continue
                        
                        is_active = _parse_boolean(row.get('is_active', True))
                        slot_name = str(row.get('slot_name', '')).strip() or None
                        
                        slot = None
                        if input_slot_number is not None and input_slot_number != 0:
                             try:
                                 slot = SlotMaster.objects.get(slot_number=input_slot_number)
                             except SlotMaster.DoesNotExist:
                                 slot = None
                        
                        if slot:
                            # Update existing
                            slot.slot_time = slot_time
                            slot.slot_name = slot_name
                            slot.is_active = is_active
                            slot.save()
                            updated_count += 1
                        else:
                            # Create new with MAX + 1
                            max_val = SlotMaster.objects.aggregate(Max('slot_number'))['slot_number__max']
                            next_val = (max_val or 0) + 1
                            SlotMaster.objects.create(
                                slot_number=next_val,
                                slot_time=slot_time,
                                slot_name=slot_name,
                                is_active=is_active
                            )
                            created_count += 1

                    except Exception as e:
                        error_count += 1
                        errors.append(f'Row {idx+2}: {str(e)}')

            
            elif master_type == 'ingredientmaster' or master_type == 'ingredient':
                for idx, row in df.iterrows():
                    try:
                        ingredient_name = str(row.get('ingredient_name', '')).strip()
                        
                        if not ingredient_name:
                            error_count += 1
                            errors.append(f'Row {idx+2}: Missing required field (ingredient_name)')
                            continue
                        
                        # Auto-generate code if not provided
                        ingredient_code = str(row.get('ingredient_code', '')).strip()
                        if not ingredient_code:
                            # Generate code
                            existing_codes = IngredientMaster.objects.filter(
                                ingredient_code__startswith='ING'
                            ).values_list('ingredient_code', flat=True)
                            
                            max_number = 0
                            for code in existing_codes:
                                try:
                                    number_part = int(code.replace('ING', ''))
                                    if number_part > max_number:
                                        max_number = number_part
                                except:
                                    continue
                            
                            next_number = max_number + 1
                            if next_number > 999:
                                error_count += 1
                                errors.append(f'Row {idx+2}: Maximum ingredient code limit reached')
                                continue
                            
                            ingredient_code = f"ING{next_number:03d}"
                        
                        is_active = _parse_boolean(row.get('is_active', True))
                        
                        ingredient, created = IngredientMaster.objects.get_or_create(
                            ingredient_code=ingredient_code,
                            defaults={
                                'ingredient_name': ingredient_name,
                                'description': str(row.get('description', '')).strip() or None,
                                'is_active': is_active,
                            }
                        )
                        
                        if not created:
                            ingredient.ingredient_name = ingredient_name
                            ingredient.description = str(row.get('description', '')).strip() or None
                            ingredient.is_active = is_active
                            ingredient.save()
                            updated_count += 1
                        else:
                            created_count += 1
                    except Exception as e:
                        error_count += 1
                        errors.append(f'Row {idx+2}: {str(e)}')
            
            elif master_type == 'treatmentroom' or master_type == 'treatment_room':
                # room_no is PK, but we might check floor/desc if needed
                # Actually user wants auto-gen for NEW.
                # So we don't require room_no in Excel, but if it's there we use it for update.
                
                for idx, row in df.iterrows():
                    try:
                        # Required if creating new? No, we auto-gen.
                        # But we need some data to create a room. Floor/Desc?
                        floor = str(row.get('floor', '')).strip() or None
                        
                        # Logic: Look for room_no
                        input_room_no = str(row.get('room_no', '')).strip()
                        room = None
                        if input_room_no:
                            try:
                                room = TreatmentRoom.objects.get(room_no=input_room_no)
                            except TreatmentRoom.DoesNotExist:
                                room = None
                        
                        if room:
                             # Update
                             room.floor = floor
                             room.description = str(row.get('description', '')).strip() or None
                             room.tr_code = str(row.get('tr_code', '')).strip() or None
                             room.is_active = _parse_boolean(row.get('is_active', True))
                             room.save()
                             updated_count += 1
                        else:
                             # Create new with TR code
                             new_no = generate_unique_code('TR', TreatmentRoom, 'room_no')
                             TreatmentRoom.objects.create(
                                 room_no=new_no,
                                 floor=floor,
                                 description=str(row.get('description', '')).strip() or None,
                                 tr_code=str(row.get('tr_code', '')).strip() or None,
                                 is_active=_parse_boolean(row.get('is_active', True))
                             )
                             created_count += 1
                             
                    except Exception as e:
                        error_count += 1
                        errors.append(f'Row {idx+2}: {str(e)}')

            
            elif master_type == 'procured':
                if 'name' not in df.columns:
                    return JsonResponse({'success': False, 'error': 'Required column missing: name'}, status=400)
                
                for idx, row in df.iterrows():
                    try:
                        name = str(row.get('name', '')).strip()
                        
                        if not name:
                            error_count += 1
                            errors.append(f'Row {idx+2}: Missing required field (name)')
                            continue
                        
                        is_active = _parse_boolean(row.get('is_active', True))
                        
                        procured, created = Procured.objects.get_or_create(
                            name=name,
                            defaults={'is_active': is_active}
                        )
                        
                        if not created:
                            procured.is_active = is_active
                            procured.save()
                            updated_count += 1
                        else:
                            created_count += 1
                    except Exception as e:
                        error_count += 1
                        errors.append(f'Row {idx+2}: {str(e)}')
            
            elif master_type == 'oilclass' or master_type == 'oil_class':
                if 'name' not in df.columns:
                    return JsonResponse({'success': False, 'error': 'Required column missing: name'}, status=400)
                
                for idx, row in df.iterrows():
                    try:
                        name = str(row.get('name', '')).strip()
                        
                        if not name:
                            error_count += 1
                            errors.append(f'Row {idx+2}: Missing required field (name)')
                            continue
                        
                        is_active = _parse_boolean(row.get('is_active', True))
                        
                        oil_class, created = OilClass.objects.get_or_create(
                            name=name,
                            defaults={'is_active': is_active}
                        )
                        
                        if not created:
                            oil_class.is_active = is_active
                            oil_class.save()
                            updated_count += 1
                        else:
                            created_count += 1
                    except Exception as e:
                        error_count += 1
                        errors.append(f'Row {idx+2}: {str(e)}')
            
            elif master_type == 'usagetype' or master_type == 'usage_type':
                if 'name' not in df.columns:
                    return JsonResponse({'success': False, 'error': 'Required column missing: name'}, status=400)
                
                for idx, row in df.iterrows():
                    try:
                        name = str(row.get('name', '')).strip()
                        
                        if not name:
                            error_count += 1
                            errors.append(f'Row {idx+2}: Missing required field (name)')
                            continue
                        
                        is_active = _parse_boolean(row.get('is_active', True))
                        
                        usage_type, created = UsageType.objects.get_or_create(
                            name=name,
                            defaults={'is_active': is_active}
                        )
                        
                        if not created:
                            usage_type.is_active = is_active
                            usage_type.save()
                            updated_count += 1
                        else:
                            created_count += 1
                    except Exception as e:
                        error_count += 1
                        errors.append(f'Row {idx+2}: {str(e)}')
            
            elif master_type == 'prescriptiontime' or master_type == 'prescription_time':
                if 'name' not in df.columns:
                    return JsonResponse({'success': False, 'error': 'Required column missing: name'}, status=400)
                
                for idx, row in df.iterrows():
                    try:
                        name = str(row.get('name', '')).strip()
                        
                        if not name:
                            error_count += 1
                            errors.append(f'Row {idx+2}: Missing required field (name)')
                            continue
                        
                        is_active = _parse_boolean(row.get('is_active', True))
                        
                        pres_time, created = PrescriptionTime.objects.get_or_create(
                            name=name,
                            defaults={'is_active': is_active}
                        )
                        
                        if not created:
                            pres_time.is_active = is_active
                            pres_time.save()
                            updated_count += 1
                        else:
                            created_count += 1
                    except Exception as e:
                        error_count += 1
                        errors.append(f'Row {idx+2}: {str(e)}')
            
            elif master_type == 'prescriptionscope' or master_type == 'prescription_scope':
                if 'name' not in df.columns:
                    return JsonResponse({'success': False, 'error': 'Required column missing: name'}, status=400)
                
                for idx, row in df.iterrows():
                    try:
                        name = str(row.get('name', '')).strip()
                        
                        if not name:
                            error_count += 1
                            errors.append(f'Row {idx+2}: Missing required field (name)')
                            continue
                        
                        is_active = _parse_boolean(row.get('is_active', True))
                        
                        pres_scope, created = PrescriptionScope.objects.get_or_create(
                            name=name,
                            defaults={'is_active': is_active}
                        )
                        
                        if not created:
                            pres_scope.is_active = is_active
                            pres_scope.save()
                            updated_count += 1
                        else:
                            created_count += 1
                    except Exception as e:
                        error_count += 1
                        errors.append(f'Row {idx+2}: {str(e)}')
            
            elif master_type == 'unitno' or master_type == 'unit_no':
                if 'unit_name' not in df.columns:
                    return JsonResponse({'success': False, 'error': 'Required column missing: unit_name'}, status=400)
                
                for idx, row in df.iterrows():
                    try:
                        unit_name = str(row.get('unit_name', '')).strip()
                        
                        if not unit_name:
                            error_count += 1
                            errors.append(f'Row {idx+2}: Missing required field (unit_name)')
                            continue
                        
                        is_active = _parse_boolean(row.get('is_active', True))
                        
                        unit, created = UnitNo.objects.get_or_create(
                            unit_name=unit_name,
                            defaults={'is_active': is_active}
                        )
                        
                        if not created:
                            unit.is_active = is_active
                            unit.save()
                            updated_count += 1
                        else:
                            created_count += 1
                    except Exception as e:
                        error_count += 1
                        errors.append(f'Row {idx+2}: {str(e)}')
            
            elif master_type == 'nursetreatment' or master_type == 'nurse_treatment':
                if 'treatment_code' not in df.columns or 'name' not in df.columns:
                    return JsonResponse({'success': False, 'error': 'Required columns missing: treatment_code, name'}, status=400)
                
                for idx, row in df.iterrows():
                    try:
                        treatment_code = str(row.get('treatment_code', '')).strip()
                        name = str(row.get('name', '')).strip()
                        
                        if not treatment_code or not name:
                            error_count += 1
                            errors.append(f'Row {idx+2}: Missing required fields (treatment_code, name)')
                            continue
                        
                        is_active = _parse_boolean(row.get('is_active', True))
                        number_of_nurse_required = _parse_int(row.get('number_of_nurse_required', 1))
                        
                        treatment, created = NurseTreatmentMaster.objects.get_or_create(
                            treatment_code=treatment_code,
                            defaults={
                                'name': name,
                                'number_of_nurse_required': number_of_nurse_required,
                                'is_active': is_active,
                            }
                        )
                        
                        if not created:
                            treatment.name = name
                            treatment.number_of_nurse_required = number_of_nurse_required
                            treatment.is_active = is_active
                            treatment.save()
                            updated_count += 1
                        else:
                            created_count += 1
                    except Exception as e:
                        error_count += 1
                        errors.append(f'Row {idx+2}: {str(e)}')
            
            elif master_type == 'nursemaster' or master_type == 'nurse_master':
                if 'nurse_name' not in df.columns:
                     return JsonResponse({'success': False, 'error': 'Required column missing: nurse_name'}, status=400)
                
                for idx, row in df.iterrows():
                    try:
                        nurse_name = str(row.get('nurse_name', '')).strip()
                        if not nurse_name:
                            error_count += 1
                            errors.append(f'Row {idx+2}: Missing required field (nurse_name)')
                            continue

                        input_code = str(row.get('nurse_code', '')).strip()
                        nurse = None
                        if input_code:
                            try:
                                nurse = NurseMaster.objects.get(nurse_code=input_code)
                            except NurseMaster.DoesNotExist:
                                nurse = None
                        
                        is_active = _parse_boolean(row.get('is_active', True))
                        designation = str(row.get('designation', '')).strip() or None
                        
                        if nurse:
                            # Update
                            nurse.nurse_name = nurse_name
                            nurse.designation = designation
                            nurse.is_active = is_active
                            nurse.save()
                            updated_count += 1
                        else:
                            # Create new
                            new_code = generate_numeric_code(NurseMaster, 'nurse_code', start_from=100)
                            NurseMaster.objects.create(
                                nurse_code=new_code,
                                nurse_name=nurse_name,
                                designation=designation,
                                is_active=is_active
                            )
                            created_count += 1

                    except Exception as e:
                        error_count += 1
                        errors.append(f'Row {idx+2}: {str(e)}')

            
            elif master_type == 'prescriptionusage' or master_type == 'prescription_usage':
                if 'name' not in df.columns:
                    return JsonResponse({'success': False, 'error': 'Required column missing: name'}, status=400)
                
                for idx, row in df.iterrows():
                    try:
                        name = str(row.get('name', '')).strip()
                        
                        if not name:
                            error_count += 1
                            errors.append(f'Row {idx+2}: Missing required field (name)')
                            continue
                        
                        is_active = _parse_boolean(row.get('is_active', True))
                        
                        pres_usage, created = PrescriptionUsage.objects.get_or_create(
                            name=name,
                            defaults={'is_active': is_active}
                        )
                        
                        if not created:
                            pres_usage.is_active = is_active
                            pres_usage.save()
                            updated_count += 1
                        else:
                            created_count += 1
                    except Exception as e:
                        error_count += 1
                        errors.append(f'Row {idx+2}: {str(e)}')
            
            elif master_type == 'therapist' or master_type == 'therapistmaster':
                if 'therapist_code' not in df.columns or 'therapist_name' not in df.columns:
                    return JsonResponse({'success': False, 'error': 'Required columns missing: therapist_code, therapist_name'}, status=400)
                
                for idx, row in df.iterrows():
                    try:
                        therapist_code = str(row.get('therapist_code', '')).strip()
                        therapist_name = str(row.get('therapist_name', '')).strip()
                        
                        if not therapist_code or not therapist_name:
                            error_count += 1
                            errors.append(f'Row {idx+2}: Missing required fields (therapist_code, therapist_name)')
                            continue
                        
                        is_active = _parse_boolean(row.get('is_active', True))
                        is_bathing = _parse_boolean(row.get('is_bathing_therapist', False))
                        therapist_gender = str(row.get('therapist_gender', '')).strip() or None
                        employee_code = str(row.get('employee_code', '')).strip() or None
                        
                        # Note: Therapist master has relationships (grade, user, etc.) that are not handled in bulk upload
                        # This is a basic implementation that only handles core fields
                        therapist, created = TherapistMaster.objects.get_or_create(
                            therapist_code=therapist_code,
                            defaults={
                                'therapist_name': therapist_name,
                                'employee_code': employee_code,
                                'therapist_gender': therapist_gender,
                                'is_bathing_therapist': is_bathing,
                                'is_active': is_active,
                            }
                        )
                        
                        if not created:
                            therapist.therapist_name = therapist_name
                            therapist.employee_code = employee_code
                            therapist.therapist_gender = therapist_gender
                            therapist.is_bathing_therapist = is_bathing
                            therapist.is_active = is_active
                            therapist.save()
                            updated_count += 1
                        else:
                            created_count += 1
                    except Exception as e:
                        error_count += 1
                        errors.append(f'Row {idx+2}: {str(e)}')
            
            else:
                return JsonResponse({'success': False, 'error': f'Unknown master type: {master_type}'}, status=400)
        
        # Prepare response
        message = f'Successfully processed {created_count + updated_count} records. Created: {created_count}, Updated: {updated_count}'
        if error_count > 0:
            message += f', Errors: {error_count}'
        
        result = {
            'success': True,
            'message': message,
            'created': created_count,
            'updated': updated_count,
            'errors': error_count,
            'total': len(df),
            'error_details': errors[:10]  # Limit to first 10 errors
        }
        
        return JsonResponse(result)
        
    except Exception as e:
        import traceback
        return JsonResponse({
            'success': False,
            'error': f'Unexpected error: {str(e)}',
            'traceback': traceback.format_exc()
        }, status=500)

def _parse_boolean(value):
    """Parse various boolean representations"""
    if value is None or pd.isna(value):
        return True
    if isinstance(value, bool):
        return value
    value_str = str(value).lower().strip()
    return value_str in ('true', '1', 'yes', 'active', 'y', 'on')

def _parse_int(value, default=0):
    """Parse integer value"""
    if value is None or pd.isna(value):
        return default
    try:
        return int(float(value))
    except:
        return default

def _parse_decimal(value):
    """Parse decimal value"""
    if value is None or pd.isna(value):
        return None
    try:
        from decimal import Decimal
        return Decimal(str(value))
    except:
        return None

@login_required
@super_admin_required
def therapist_master_delete(request, pk):
    """Delete a therapist - function-based view to handle GET requests properly"""
    therapist = get_object_or_404(TherapistMaster, pk=pk)
    therapist_name = therapist.therapist_name
    therapist_code = therapist.therapist_code
    therapist.delete()
    _log_master_audit(request, 'delete', 'TherapistMaster', str(pk), f'Deleted TherapistMaster: {therapist_name} ({therapist_code})')
    messages.success(request, f'Therapist "{therapist_name}" ({therapist_code}) deleted successfully.')
    return redirect('therapist_master_manage')


@login_required
@super_admin_required
def therapist_master_enable(request, pk):
    """Enable a therapist master record."""
    therapist = get_object_or_404(TherapistMaster, pk=pk)
    therapist.is_active = True
    therapist.save()
    _log_master_audit(request, 'update', 'TherapistMaster', str(pk), f'Enabled TherapistMaster: {therapist.therapist_name} ({therapist.therapist_code})')
    messages.success(request, f'Therapist "{therapist.therapist_name}" ({therapist.therapist_code}) enabled successfully.')
    return redirect('therapist_master_manage')


@login_required
@super_admin_required
def therapist_master_disable(request, pk):
    """Disable a therapist master record."""
    therapist = get_object_or_404(TherapistMaster, pk=pk)
    therapist.is_active = False
    therapist.save()
    _log_master_audit(request, 'update', 'TherapistMaster', str(pk), f'Disabled TherapistMaster: {therapist.therapist_name} ({therapist.therapist_code})')
    messages.success(request, f'Therapist "{therapist.therapist_name}" ({therapist.therapist_code}) disabled successfully.')
    return redirect('therapist_master_manage')

# Therapist Grade CRUD Views
class TherapistGradeManagementView(ListView):
    model = TherapistGrade
    template_name = 'masters/therapistgrade_manage.html'
    context_object_name = 'grade_list'
    paginate_by = 20

    def get_queryset(self):
        queryset = TherapistGrade.objects.all()
        
        # Search filter
        search = self.request.GET.get('search', '').strip()
        if search:
            queryset = queryset.filter(
                db_models.Q(name__icontains=search)
            )
        
        # Status filter
        status = self.request.GET.get('status', '')
        if status == 'active':
            queryset = queryset.filter(is_active=True)
        elif status == 'inactive':
            queryset = queryset.filter(is_active=False)
        # If status is empty or 'all', show all grades (no filter applied)
        
        return queryset.order_by('display_order', 'name')
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['search'] = self.request.GET.get('search', '')
        context['status_filter'] = self.request.GET.get('status', '')
        return context

class TherapistGradeCreateView(CreateView):
    model = TherapistGrade
    fields = ['name', 'display_order', 'is_active']
    template_name = 'masters/therapistgrade_form.html'
    success_url = reverse_lazy('therapistgrade_manage')
    
    def form_valid(self, form):
        response = super().form_valid(form)
        _log_master_audit(self.request, 'create', 'TherapistGrade', str(self.object.pk), f'Created TherapistGrade: {self.object.name}')
        messages.success(self.request, 'Therapist grade created successfully.')
        return response

class TherapistGradeUpdateView(UpdateView):
    model = TherapistGrade
    fields = ['name', 'display_order', 'is_active']
    template_name = 'masters/therapistgrade_form.html'
    success_url = reverse_lazy('therapistgrade_manage')
    
    def form_valid(self, form):
        _log_master_audit(self.request, 'update', 'TherapistGrade', str(self.object.pk), f'Updated TherapistGrade: {self.object.name}')
        messages.success(self.request, 'Therapist grade updated successfully.')
        return super().form_valid(form)

@login_required
@super_admin_required
def therapist_grade_delete(request, pk):
    """Delete a therapist grade - function-based view"""
    grade = get_object_or_404(TherapistGrade, pk=pk)
    grade_name = grade.name
    # Check if any therapists are using this grade
    if TherapistMaster.objects.filter(therapist_grade=grade).exists():
        messages.error(request, f'Cannot delete grade "{grade_name}" because it is assigned to one or more therapists.')
        return redirect('therapistgrade_manage')
    grade.delete()
    _log_master_audit(request, 'delete', 'TherapistGrade', str(pk), f'Deleted TherapistGrade: {grade_name}')
    messages.success(request, f'Therapist grade "{grade_name}" deleted successfully.')
    return redirect('therapistgrade_manage')


@login_required
@super_admin_required
def therapist_grade_enable(request, pk):
    """Enable a therapist grade record."""
    grade = get_object_or_404(TherapistGrade, pk=pk)
    grade.is_active = True
    grade.save()
    _log_master_audit(request, 'update', 'TherapistGrade', str(pk), f'Enabled TherapistGrade: {grade.name}')
    messages.success(request, f'Therapist grade "{grade.name}" enabled successfully.')
    return redirect('therapistgrade_manage')


@login_required
@super_admin_required
def therapist_grade_disable(request, pk):
    """Disable a therapist grade record."""
    grade = get_object_or_404(TherapistGrade, pk=pk)
    grade.is_active = False
    grade.save()
    _log_master_audit(request, 'update', 'TherapistGrade', str(pk), f'Disabled TherapistGrade: {grade.name}')
    messages.success(request, f'Therapist grade "{grade.name}" disabled successfully.')
    return redirect('therapistgrade_manage')

# Therapist Management API endpoints
@login_required
@super_admin_required
@require_http_methods(["POST"])
@csrf_exempt
def create_therapist(request):
    """API endpoint to create therapist (Super Admin only) - Creates user account and profile"""
    try:
        data = json.loads(request.body)
        
        therapist_code = data['therapist_code']
        therapist_name = data['therapist_name']
        therapist_grade = data.get('therapist_grade', 'Junior')
        employee_code = data['employee_code']
        is_bathing_therapist = data.get('is_bathing_therapist', True)
        is_active = data.get('is_active', True)
        
        # Check if therapist already exists
        if TherapistMaster.objects.filter(therapist_code=therapist_code).exists():
            return JsonResponse({'error': f'Therapist with code "{therapist_code}" already exists'}, status=400)
        
        # Generate username from therapist name or code
        username = data.get('username') or therapist_code.lower().replace(' ', '_')
        email = data.get('email') or f"{username}@avcri.local"
        password = data.get('password') or f"{therapist_code}@{employee_code}"  # Default password
        
        # Create user account
        from authentication.models import CustomUser, TherapistProfile
        from django.contrib.auth import get_user_model
        User = get_user_model()
        
        # Check if username already exists
        if User.objects.filter(username=username).exists():
            username = f"{username}_{therapist_code}"  # Append code if duplicate
        
        # Split name into first and last name
        name_parts = therapist_name.split(maxsplit=1)
        first_name = name_parts[0] if name_parts else therapist_name
        last_name = name_parts[1] if len(name_parts) > 1 else ''
        
        user = User.objects.create_user(
            username=username,
            email=email,
            password=password,
            first_name=first_name,
            last_name=last_name,
            role='therapist',
            mobile=data.get('mobile', ''),
            gender=data.get('gender', 'male')
        )
        
        # Create therapist profile
        therapist_profile = TherapistProfile.objects.create(
            user=user,
            therapist_code=therapist_code,
            employee_code=employee_code,
            therapist_grade=therapist_grade,
            is_bathing_therapist=is_bathing_therapist,
            department=data.get('department', ''),
            specialization=data.get('specialization', 'General Therapy')
        )
        
        # Create therapist master record
        therapist = TherapistMaster.objects.create(
            user=user,
            therapist_profile=therapist_profile,
            therapist_code=therapist_code,
            therapist_name=therapist_name,
            therapist_grade=therapist_grade,
            employee_code=employee_code,
            is_bathing_therapist=is_bathing_therapist,
            is_active=is_active
        )
        _log_master_audit(request, 'create', 'TherapistMaster', str(therapist.id), f'Created TherapistMaster (API): {therapist_name} ({therapist_code})')
        
        return JsonResponse({
            'success': True, 
            'id': therapist.id,
            'user_id': user.id,
            'username': username,
            'message': f'Therapist user account created. Default password: {password}'
        })
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({'error': str(e)}, status=500)

@login_required
@super_admin_required
@require_http_methods(["POST"])
@csrf_exempt
def create_therapist_leave(request):
    """API endpoint to create therapist leave (Super Admin only)"""
    try:
        data = json.loads(request.body)
        
        therapist = get_object_or_404(TherapistMaster, therapist_code=data['therapist_code'])
        
        leave = TherapistLeaveMaster.objects.create(
            therapist=therapist,
            leave_day=data['leave_day'],
            is_active=data.get('is_active', True)
        )
        _log_master_audit(request, 'create', 'TherapistLeaveMaster', str(leave.id), f'Created TherapistLeaveMaster (API): therapist {therapist.therapist_code}, day {leave.leave_day}')
        
        return JsonResponse({'success': True, 'id': leave.id})
        
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)

@login_required
@super_admin_required
@require_http_methods(["PUT"])
@csrf_exempt
def update_therapist(request, therapist_id):
    """API endpoint to update therapist (Super Admin only)"""
    try:
        data = json.loads(request.body)
        therapist = get_object_or_404(TherapistMaster, id=therapist_id)
        
        therapist.therapist_name = data.get('therapist_name', therapist.therapist_name)
        therapist.therapist_grade = data.get('therapist_grade', therapist.therapist_grade)
        therapist.employee_code = data.get('employee_code', therapist.employee_code)
        therapist.is_bathing_therapist = data.get('is_bathing_therapist', therapist.is_bathing_therapist)
        therapist.is_active = data.get('is_active', therapist.is_active)
        therapist.save()
        _log_master_audit(request, 'update', 'TherapistMaster', str(therapist_id), f'Updated TherapistMaster (API): {therapist.therapist_name} ({therapist.therapist_code})')
        
        return JsonResponse({'success': True})
        
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)

@login_required
@super_admin_required
@require_http_methods(["DELETE"])
@csrf_exempt
def delete_therapist(request, therapist_id):
    """API endpoint to delete therapist (Super Admin only)"""
    try:
        therapist = get_object_or_404(TherapistMaster, id=therapist_id)
        name, code = therapist.therapist_name, therapist.therapist_code
        therapist.delete()
        _log_master_audit(request, 'delete', 'TherapistMaster', str(therapist_id), f'Deleted TherapistMaster (API): {name} ({code})')
        return JsonResponse({'success': True})
        
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)

@login_required
@super_admin_required
@require_http_methods(["DELETE"])
@csrf_exempt
def delete_therapist_leave(request, leave_id):
    """API endpoint to delete therapist leave (Super Admin only)"""
    try:
        leave = get_object_or_404(TherapistLeaveMaster, id=leave_id)
        info = f'therapist {leave.therapist.therapist_code}, day {leave.leave_day}'
        leave.delete()
        _log_master_audit(request, 'delete', 'TherapistLeaveMaster', str(leave_id), f'Deleted TherapistLeaveMaster (API): {info}')
        return JsonResponse({'success': True})
        
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)


# ── Room Name Master CRUD ──

@method_decorator(login_required, name='dispatch')
@method_decorator(super_admin_required, name='dispatch')
class RoomNameManagementView(ListView):
    model = RoomName
    template_name = 'masters/roomname_manage.html'
    context_object_name = 'roomname_list'
    paginate_by = 20

    def get_queryset(self):
        queryset = RoomName.objects.all()
        search = self.request.GET.get('search', '').strip()
        if search:
            queryset = queryset.filter(name__icontains=search)
        status = self.request.GET.get('status', '').strip().lower()
        if status == 'active':
            queryset = queryset.filter(is_active=True)
        elif status == 'inactive':
            queryset = queryset.filter(is_active=False)
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['search'] = self.request.GET.get('search', '')
        context['status_filter'] = self.request.GET.get('status', '')
        context['master_type'] = 'roomname'
        return context

    def post(self, request, *args, **kwargs):
        name = request.POST.get('name', '').strip()
        is_active = request.POST.get('is_active') == 'on'
        if not name:
            messages.error(request, 'Room name is required.')
            return redirect('roomname_manage')
        if RoomName.objects.filter(name__iexact=name).exists():
            messages.error(request, f'Room name "{name}" already exists.')
            return redirect('roomname_manage')
        obj = RoomName.objects.create(name=name, is_active=is_active)
        _log_master_audit(request, 'create', 'RoomName', str(obj.pk), f'Created RoomName: {obj.name}')
        messages.success(request, f'Room name "{obj.name}" created successfully.')
        return redirect('roomname_manage')


@method_decorator(login_required, name='dispatch')
@method_decorator(super_admin_required, name='dispatch')
class RoomNameUpdateView(UpdateView):
    model = RoomName
    fields = ['name', 'is_active']
    template_name = 'masters/roomname_form.html'
    success_url = reverse_lazy('roomname_manage')

    def form_valid(self, form):
        _log_master_audit(self.request, 'update', 'RoomName', str(self.object.pk), f'Updated RoomName: {self.object.name}')
        messages.success(self.request, 'Room name updated successfully.')
        return super().form_valid(form)


@login_required
@super_admin_required
def roomname_delete(request, pk):
    obj = get_object_or_404(RoomName, pk=pk)
    name = obj.name
    obj.delete()
    _log_master_audit(request, 'delete', 'RoomName', str(pk), f'Deleted RoomName: {name}')
    messages.success(request, f'Room name "{name}" deleted successfully.')
    return redirect('roomname_manage')


@login_required
@super_admin_required
def roomname_enable(request, pk):
    obj = get_object_or_404(RoomName, pk=pk)
    obj.is_active = True
    obj.save()
    _log_master_audit(request, 'update', 'RoomName', str(pk), f'Enabled RoomName: {obj.name}')
    messages.success(request, f'Room name "{obj.name}" enabled successfully.')
    return redirect('roomname_manage')


@login_required
@super_admin_required
def roomname_disable(request, pk):
    obj = get_object_or_404(RoomName, pk=pk)
    obj.is_active = False
    obj.save()
    _log_master_audit(request, 'update', 'RoomName', str(pk), f'Disabled RoomName: {obj.name}')
    messages.success(request, f'Room name "{obj.name}" disabled successfully.')
    return redirect('roomname_manage')


# ── Floor Master CRUD ──

@method_decorator(login_required, name='dispatch')
@method_decorator(super_admin_required, name='dispatch')
class FloorManagementView(ListView):
    model = Floor
    template_name = 'masters/floor_manage.html'
    context_object_name = 'floor_list'
    paginate_by = 20

    def get_queryset(self):
        queryset = Floor.objects.all()
        search = self.request.GET.get('search', '').strip()
        if search:
            queryset = queryset.filter(name__icontains=search)
        status = self.request.GET.get('status', '').strip().lower()
        if status == 'active':
            queryset = queryset.filter(is_active=True)
        elif status == 'inactive':
            queryset = queryset.filter(is_active=False)
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['search'] = self.request.GET.get('search', '')
        context['status_filter'] = self.request.GET.get('status', '')
        context['master_type'] = 'floor'
        return context

    def post(self, request, *args, **kwargs):
        name = request.POST.get('name', '').strip()
        is_active = request.POST.get('is_active') == 'on'
        if not name:
            messages.error(request, 'Floor name is required.')
            return redirect('floor_manage')
        if Floor.objects.filter(name__iexact=name).exists():
            messages.error(request, f'Floor "{name}" already exists.')
            return redirect('floor_manage')
        obj = Floor.objects.create(name=name, is_active=is_active)
        _log_master_audit(request, 'create', 'Floor', str(obj.pk), f'Created Floor: {obj.name}')
        messages.success(request, f'Floor "{obj.name}" created successfully.')
        return redirect('floor_manage')


@method_decorator(login_required, name='dispatch')
@method_decorator(super_admin_required, name='dispatch')
class FloorUpdateView(UpdateView):
    model = Floor
    fields = ['name', 'is_active']
    template_name = 'masters/floor_form.html'
    success_url = reverse_lazy('floor_manage')

    def form_valid(self, form):
        _log_master_audit(self.request, 'update', 'Floor', str(self.object.pk), f'Updated Floor: {self.object.name}')
        messages.success(self.request, 'Floor updated successfully.')
        return super().form_valid(form)


@login_required
@super_admin_required
def floor_delete(request, pk):
    obj = get_object_or_404(Floor, pk=pk)
    name = obj.name
    obj.delete()
    _log_master_audit(request, 'delete', 'Floor', str(pk), f'Deleted Floor: {name}')
    messages.success(request, f'Floor "{name}" deleted successfully.')
    return redirect('floor_manage')


@login_required
@super_admin_required
def floor_enable(request, pk):
    obj = get_object_or_404(Floor, pk=pk)
    obj.is_active = True
    obj.save()
    _log_master_audit(request, 'update', 'Floor', str(pk), f'Enabled Floor: {obj.name}')
    messages.success(request, f'Floor "{obj.name}" enabled successfully.')
    return redirect('floor_manage')


@login_required
@super_admin_required
def floor_disable(request, pk):
    obj = get_object_or_404(Floor, pk=pk)
    obj.is_active = False
    obj.save()
    _log_master_audit(request, 'update', 'Floor', str(pk), f'Disabled Floor: {obj.name}')
    messages.success(request, f'Floor "{obj.name}" disabled successfully.')
    return redirect('floor_manage')